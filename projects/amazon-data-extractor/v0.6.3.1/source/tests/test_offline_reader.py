from __future__ import annotations

import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from amazon_data_extractor.offline_reader import OfflineWorkbookError, read_offline_extraction

class OfflineReaderTests(unittest.TestCase):
    def test_reads_records_and_accepted_asins(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "amazon_extraction.xlsx"
            workbook = Workbook()
            data = workbook.active
            data.title = "Amazon Data"
            data.append(["input_asin", "asin", "title", "features"])
            data.append(["B000000001", "B000000001", "Example", json.dumps(["One", "Two"])])
            audit = workbook.create_sheet("Input Audit")
            audit.append(["Source Row", "Original Value", "Normalized ASIN", "Status", "Message"])
            audit.append([2, "B000000001", "B000000001", "ACCEPTED", ""])
            audit.append([3, "B000000001", "B000000001", "DUPLICATE", "Earlier row"])
            workbook.save(path)

            result = read_offline_extraction(path)
            self.assertEqual(result.requested_asins, ["B000000001"])
            self.assertEqual(len(result.records), 1)
            self.assertEqual(result.records[0]["title"], "Example")

    def test_rejects_unrelated_workbook(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "other.xlsx"
            workbook = Workbook()
            workbook.save(path)
            with self.assertRaisesRegex(OfflineWorkbookError, "Input Audit"):
                read_offline_extraction(path)


if __name__ == "__main__":
    unittest.main()
