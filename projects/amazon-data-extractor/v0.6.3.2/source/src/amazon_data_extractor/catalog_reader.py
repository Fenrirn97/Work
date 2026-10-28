"""Read Amazon catalog rows without assuming a fixed worksheet name."""

from __future__ import annotations

from dataclasses import dataclass, replace
from pathlib import Path
import re
import unicodedata
from typing import Iterable

from openpyxl import load_workbook


ASIN_PATTERN = re.compile(r"^[A-Z0-9]{10}$", re.IGNORECASE)


@dataclass(frozen=True)
class CatalogProduct:
    asin: str
    sku: str
    product_name: str
    bullets: list[str]
    identifier_type: str = ""
    identifier_value: str = ""
    asin_source: str = "Catalog ASIN"
    source_row: int = 0


@dataclass(frozen=True)
class CatalogData:
    products_by_asin: dict[str, list[CatalogProduct]]
    products_by_sku: dict[str, list[CatalogProduct]]
    products: list[CatalogProduct]
    sheet_name: str
    record_count: int


def clean_asin(value: object) -> str:
    """Return a valid uppercase ASIN or an empty string."""
    text = str(value or "").strip().upper()
    return text if ASIN_PATTERN.fullmatch(text) else ""


def clean_text(value: object) -> str:
    return str(value or "").strip()


def sku_key(value: object) -> str:
    return clean_text(value).casefold()


def normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", clean_text(value)).casefold()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _first_column(columns: dict[str, int], aliases: Iterable[str]) -> int | None:
    for alias in aliases:
        if alias in columns:
            return columns[alias]
    return None


def find_catalog_columns(worksheet) -> tuple[int, dict[str, object]] | None:
    """Find the human-readable header row in supported Amazon reports."""
    best_match = None
    best_score = -1

    for row_number, row in enumerate(
        worksheet.iter_rows(min_row=1, max_row=25, values_only=True),
        start=1,
    ):
        headers = [normalize_header(value) for value in row]
        first_column: dict[str, int] = {}
        for index, header in enumerate(headers):
            if header and header not in first_column:
                first_column[header] = index

        sku_column = _first_column(first_column, ("sku", "seller sku"))
        if sku_column is None:
            continue

        numbered_bullets: dict[int, int] = {}
        for index, header in enumerate(headers):
            match = re.fullmatch(r"bullet ?point ([1-5])", header)
            if match:
                numbered_bullets[int(match.group(1))] = index
        duplicate_bullets = [
            index for index, header in enumerate(headers) if header == "bullet point"
        ]
        for number, index in enumerate(duplicate_bullets[:5], start=1):
            numbered_bullets.setdefault(number, index)
        if not numbered_bullets:
            continue

        asin_column = first_column.get("asin")
        product_id_type_column = _first_column(
            first_column,
            (
                "product id type",
                "external product id type",
                "product identifier type",
                "identifier type",
            ),
        )
        product_id_column = _first_column(
            first_column,
            (
                "product id",
                "external product id",
                "product identifier",
                "identifier number",
            ),
        )
        gtin_exemption_column = _first_column(
            first_column,
            ("gtin exemption reason", "gtin exempt", "gtin exemption"),
        )
        name_column = _first_column(
            first_column,
            ("title", "product name", "item name", "item name source field"),
        )

        has_identifier = asin_column is not None or (
            product_id_type_column is not None and product_id_column is not None
        )
        score = (
            len(numbered_bullets) * 3
            + (8 if has_identifier else 0)
            + (2 if name_column is not None else 0)
        )
        if score > best_score:
            best_score = score
            best_match = (
                row_number,
                {
                    "asin": asin_column,
                    "product_id_type": product_id_type_column,
                    "product_id": product_id_column,
                    "gtin_exemption": gtin_exemption_column,
                    "sku": sku_column,
                    "name": name_column,
                    "bullets": numbered_bullets,
                },
            )

    return best_match


def first_data_row(worksheet, header_row: int, columns: dict[str, object]) -> int:
    """Skip the machine-field and example rows in Amazon Template sheets."""
    data_row = header_row + 1
    if worksheet.title != "Template":
        return data_row

    probe = next(
        worksheet.iter_rows(min_row=data_row, max_row=data_row, values_only=True),
        None,
    )
    if not probe:
        return data_row

    sku_column = columns["sku"]
    sku_text = clean_text(probe[sku_column]) if sku_column < len(probe) else ""
    if "#1.value" in sku_text or "contribution_sku" in sku_text.casefold():
        return header_row + 3
    return data_row


def _row_value(row: tuple, column: int | None) -> str:
    return clean_text(row[column]) if column is not None and column < len(row) else ""


def _identifier_from_row(row: tuple, columns: dict[str, object]) -> tuple[str, str]:
    identifier_type = _row_value(row, columns["product_id_type"])
    identifier_value = _row_value(row, columns["product_id"])
    if not identifier_value:
        exemption = _row_value(row, columns["gtin_exemption"])
        if exemption:
            identifier_type = identifier_type or "GTIN Exempt"
            identifier_value = exemption
    return identifier_type, identifier_value


def _asin_from_row(
    row: tuple,
    columns: dict[str, object],
    identifier_type: str,
    identifier_value: str,
) -> str:
    asin = clean_asin(_row_value(row, columns["asin"]))
    if asin:
        return asin
    if identifier_type.casefold() == "asin":
        return clean_asin(identifier_value)
    return ""


def build_catalog_data(
    products: list[CatalogProduct],
    *,
    sheet_name: str,
) -> CatalogData:
    by_asin: dict[str, list[CatalogProduct]] = {}
    by_sku: dict[str, list[CatalogProduct]] = {}
    for product in products:
        if product.asin:
            by_asin.setdefault(product.asin, []).append(product)
        if product.sku:
            by_sku.setdefault(sku_key(product.sku), []).append(product)
    return CatalogData(by_asin, by_sku, products, sheet_name, len(products))


def with_product_asins(
    catalog: CatalogData,
    replacements: dict[int, tuple[str, str]],
) -> CatalogData:
    """Return a re-indexed catalog after resolving selected blank ASINs."""
    products = []
    for product in catalog.products:
        replacement = replacements.get(product.source_row)
        if product.asin or not replacement:
            products.append(product)
            continue
        asin, source = replacement
        products.append(replace(product, asin=asin, asin_source=source))
    return build_catalog_data(products, sheet_name=catalog.sheet_name)


def load_catalog(path: str | Path) -> CatalogData:
    """Load ASIN, identifier, SKU, name, and Bullet Point 1-5 fields."""
    path = Path(path)
    workbook = load_workbook(
        path,
        read_only=True,
        data_only=True,
        keep_vba=path.suffix.lower() == ".xlsm",
    )

    try:
        preferred_sheets = [
            name
            for name in ("Template", "Listing Content", "Products")
            if name in workbook.sheetnames
        ]
        preferred_sheets += [
            name for name in workbook.sheetnames if name not in preferred_sheets
        ]

        selected = None
        for sheet_name in preferred_sheets:
            worksheet = workbook[sheet_name]
            columns = find_catalog_columns(worksheet)
            if columns:
                selected = (worksheet, columns)
                break

        if selected is None:
            found = ", ".join(workbook.sheetnames) or "(none)"
            raise ValueError(
                "Category Listings Report was not recognized. It needs SKU and "
                "Bullet Point columns; ASIN or Product ID fields are recommended. "
                f"Worksheets found: {found}."
            )

        worksheet, (header_row, columns) = selected
        start_row = first_data_row(worksheet, header_row, columns)
        products: list[CatalogProduct] = []

        for source_row, row in enumerate(
            worksheet.iter_rows(min_row=start_row, values_only=True),
            start=start_row,
        ):
            sku = _row_value(row, columns["sku"])
            if not sku:
                continue

            identifier_type, identifier_value = _identifier_from_row(row, columns)
            asin = _asin_from_row(row, columns, identifier_type, identifier_value)
            if asin and not identifier_type:
                identifier_type = "ASIN"
            if asin and not identifier_value:
                identifier_value = asin

            bullets = [
                _row_value(row, columns["bullets"].get(number))
                for number in range(1, 6)
            ]
            products.append(
                CatalogProduct(
                    asin=asin,
                    sku=sku,
                    product_name=_row_value(row, columns["name"]),
                    bullets=bullets,
                    identifier_type=identifier_type,
                    identifier_value=identifier_value,
                    asin_source="Catalog ASIN" if asin else "",
                    source_row=source_row,
                )
            )

        return build_catalog_data(products, sheet_name=worksheet.title)
    finally:
        workbook.close()
