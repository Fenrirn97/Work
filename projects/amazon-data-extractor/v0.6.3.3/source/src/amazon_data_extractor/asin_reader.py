"""Read ASIN lists and Amazon Category Listings Reports without modifying them."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata

from openpyxl import load_workbook

from .catalog_reader import first_data_row

ASIN_PATTERN = re.compile(r"^[A-Z0-9]{10}$")
AMAZON_URL_PATTERNS = (
    re.compile(r"/(?:dp|gp/product|product)/([A-Z0-9]{10})(?:[/?#]|$)", re.I),
    re.compile(r"[?&]asin=([A-Z0-9]{10})(?:[&#]|$)", re.I),
)
DEFAULT_COLUMN_ALIASES = ("child asin", "asin", "amazon asin", "product asin", "asin1", "asin 1")
MAX_HEADER_ROWS = 50
MAX_HEADER_COLUMNS = 2048
HELPER_SHEETS = {
    "instructions", "images", "data definitions", "valid values", "dropdown lists",
    "conditions list", "attributeptdmap", "changes to the template",
}
ALIASES = {
    "sku": ("sku", "seller sku", "item sku", "contribution sku"),
    "name": ("item name", "product name", "title"),
    "product_id": ("product id", "external product id", "product identifier", "identifier number", "product id value"),
    "product_id_type": ("product id type", "external product id type", "product identifier type", "identifier type"),
    "parentage": ("parentage level", "parentage", "parent child", "parent child relationship"),
    "status": ("status", "listing status"),
}


class InputWorkbookError(ValueError):
    """The workbook cannot provide a usable ASIN input."""


@dataclass(frozen=True)
class AsinAuditRow:
    source_row: int
    original_value: str
    normalized_asin: str
    status: str
    message: str
    source_sku: str = ""
    identifier_type: str = ""
    source_field: str = ""


@dataclass(frozen=True)
class AsinReadResult:
    asins: list[str]
    audit_rows: list[AsinAuditRow]
    sheet_name: str
    column_name: str
    header_row: int = 1
    data_start_row: int = 2
    input_kind: str = "ASIN workbook"

    @property
    def invalid_count(self) -> int:
        return sum(row.status == "INVALID" for row in self.audit_rows)

    @property
    def duplicate_count(self) -> int:
        return sum(row.status == "DUPLICATE" for row in self.audit_rows)

    @property
    def parent_count(self) -> int:
        return sum(row.status == "SKIPPED_PARENT" for row in self.audit_rows)


def _normalize_header(value: object) -> str:
    text = unicodedata.normalize("NFKC", str(value or "").strip()).casefold()
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def _header_role(value: object) -> str:
    text = str(value or "").strip().casefold()
    # Machine-field rows occur beneath Amazon's human-readable headings.
    if text.startswith(("contribution_sku", "item_sku")):
        return "sku"
    if text == "amzn1.volt.ca.product_id_type":
        return "product id type"
    if text == "amzn1.volt.ca.product_id_value":
        return "product id"
    if text.startswith("externally_assigned_product_identifier"):
        return "product id type" if text.endswith(".type") else "product id"
    if text.startswith(("merchant_suggested_asin", "asin[")):
        return "asin"
    if text.startswith("parentage_level"):
        return "parentage level"
    return _normalize_header(value)


def extract_asin(value: object) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if ASIN_PATTERN.fullmatch(text.upper()):
        return text.upper()
    for pattern in AMAZON_URL_PATTERNS:
        match = pattern.search(text)
        if match:
            return match.group(1).upper()
    return None


def _first_column(headers: list[object], aliases: tuple[str, ...]) -> int | None:
    roles = [_header_role(value) for value in headers]
    for alias in aliases:
        if alias in roles:
            return roles.index(alias)
    return None


def _row_value(row: tuple, index: int | None) -> str:
    return str(row[index]).strip() if index is not None and index < len(row) and row[index] is not None else ""


def _find_layout(worksheet, column_name: str | None, header_row: int | None):
    minimum = header_row or 1
    maximum = header_row or MAX_HEADER_ROWS
    for number, row in enumerate(worksheet.iter_rows(
        min_row=minimum, max_row=maximum, max_col=min(worksheet.max_column or 1, MAX_HEADER_COLUMNS),
        values_only=True,
    ), start=minimum):
        headers = list(row)
        columns = {role: _first_column(headers, aliases) for role, aliases in ALIASES.items()}
        direct_aliases = (_normalize_header(column_name),) if column_name else DEFAULT_COLUMN_ALIASES
        columns["asin"] = _first_column(headers, direct_aliases)
        paired = columns["product_id"] is not None and columns["product_id_type"] is not None
        selected_id = column_name and columns["asin"] is not None and columns["asin"] == columns["product_id"]
        if selected_id:
            # A manual Product Id override must still respect its type.
            columns["asin"] = None
        if columns["asin"] is None and ((column_name and not selected_id) or not paired):
            continue
        category = paired or columns["parentage"] is not None or (
            worksheet.title.casefold() == "template" and number > 1
        )
        return {"headers": headers, "columns": columns, "header_row": number, "category": category}
    return None


def _template_row(row: tuple, layout: dict) -> bool:
    if not any(str(value).strip() for value in row if value is not None):
        return True
    columns = layout["columns"]
    if (layout["category"] and columns["sku"] is not None
            and not _row_value(row, columns["sku"])
            and not _row_value(row, columns["asin"])
            and not _row_value(row, columns["product_id"])):
        # Amazon's example/default row can contain parentage/default choices
        # but neither a listing SKU nor an actual product identifier.
        return True
    matches = 0
    for role, index in columns.items():
        if index is None:
            continue
        aliases = DEFAULT_COLUMN_ALIASES if role == "asin" else ALIASES[role]
        if _header_role(_row_value(row, index)) in aliases:
            matches += 1
    if matches >= 2:
        return True
    # Explicit example labels are metadata, not an actual listing.
    return layout["category"] and _row_value(row, columns["status"]).casefold() in {"example", "example row"}


def _row_identifier(row: tuple, layout: dict):
    columns, headers = layout["columns"], layout["headers"]
    direct = _row_value(row, columns["asin"])
    identifier = _row_value(row, columns["product_id"])
    kind = _row_value(row, columns["product_id_type"])
    direct_asin = extract_asin(direct)
    typed_asin = extract_asin(identifier) if kind.casefold() == "asin" else None
    direct_field = str(headers[columns["asin"]]).strip() if columns["asin"] is not None else ""
    id_field = str(headers[columns["product_id"]]).strip() if columns["product_id"] is not None else ""
    if direct_asin:
        if typed_asin and typed_asin != direct_asin:
            return direct, None, "ASIN", direct_field, "Conflicting ASIN and ASIN-typed Product Id; row skipped"
        return direct, direct_asin, "ASIN", direct_field, ""
    if kind.casefold() == "asin":
        return identifier, typed_asin, kind, id_field, (
            "" if typed_asin else "ASIN-typed Product Id is blank or invalid"
        )
    if identifier or kind:
        return identifier, None, kind, id_field, (
            f"Product Id Type '{kind or 'blank'}' is not ASIN; provide the actual ASIN, not a barcode"
        )
    return direct, None, kind, direct_field or id_field, (
        "ASIN cell is blank" if not direct else "Expected a 10-character ASIN or Amazon product URL"
    )


def read_asins(
    workbook_path: str | Path, *, sheet_name: str | None = None,
    column_name: str | None = None, header_row: int | None = None,
) -> AsinReadResult:
    """Auto-detect a data sheet/header and return only unique usable child ASINs."""
    path = Path(workbook_path)
    if not path.is_file():
        raise InputWorkbookError(f"Excel file not found: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise InputWorkbookError("Use an .xlsx or .xlsm workbook, including Amazon's Category Listings Report.")
    if header_row is not None and header_row < 1:
        raise InputWorkbookError("Header row must be 1 or greater.")
    try:
        # No macro execution and no save/write operation is performed.
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise InputWorkbookError(f"Could not open the Excel workbook: {exc}") from exc
    try:
        if sheet_name:
            if sheet_name not in workbook.sheetnames:
                raise InputWorkbookError(f"Worksheet '{sheet_name}' was not found.")
            candidates = [workbook[sheet_name]]
        else:
            candidates = []
            ordered = [sheet for sheet in workbook if sheet.title.casefold() == "template"]
            ordered += [workbook.active] + list(workbook)
            for sheet in ordered:
                if sheet not in candidates and sheet.sheet_state == "visible" and _normalize_header(sheet.title) not in HELPER_SHEETS:
                    candidates.append(sheet)
        selected = None
        for sheet in candidates:
            layout = _find_layout(sheet, column_name, header_row)
            if layout is not None:
                selected = sheet, layout
                break
        if selected is None:
            raise InputWorkbookError(
                "No ASIN column or Product Id / Product Id Type pair was found in the data sheets. "
                f"Auto-detection checks the first {MAX_HEADER_ROWS} rows and {MAX_HEADER_COLUMNS} columns. "
                "UPC, EAN and GTIN values cannot be used as ASINs."
            )
        worksheet, layout = selected
        columns, headers = layout["columns"], layout["headers"]
        accepted, audit, seen = [], [], set()
        data_start = layout["header_row"] + 1
        if columns["sku"] is not None:
            # Standard Amazon Template sheets reserve an example row after
            # their machine-field row; share the existing catalog reader rule.
            data_start = first_data_row(worksheet, layout["header_row"], columns)
        first_listing = True
        for number, row in enumerate(worksheet.iter_rows(
            min_row=data_start, max_col=min(worksheet.max_column or len(headers), MAX_HEADER_COLUMNS),
            values_only=True,
        ), start=data_start):
            if _template_row(row, layout):
                if first_listing:
                    data_start = number + 1
                continue
            first_listing = False
            sku = _row_value(row, columns["sku"])
            original, asin, kind, field, message = _row_identifier(row, layout)
            if _row_value(row, columns["parentage"]).casefold() == "parent":
                status, message = "SKIPPED_PARENT", "Parent listing excluded; only child/product ASINs are submitted"
            elif asin is None:
                status = "INVALID"
            elif asin in seen:
                status, message = "DUPLICATE", "Already accepted from an earlier row"
            else:
                status = "ACCEPTED"
                seen.add(asin)
                accepted.append(asin)
            audit.append(AsinAuditRow(number, original, asin or "", status, message, sku, kind, field))
        index = columns["asin"] if columns["asin"] is not None else columns["product_id"]
        return AsinReadResult(accepted, audit, worksheet.title, str(headers[index]).strip(),
                              layout["header_row"], data_start,
                              "Category Listings Report" if layout["category"] else "ASIN workbook")
    finally:
        workbook.close()
