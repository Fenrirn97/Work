"""Synthetic Category Listings input fixtures; no private catalog is bundled."""

from __future__ import annotations

import contextlib
import hashlib
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import Mock, patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from openpyxl import Workbook, load_workbook

from app import create_app
from amazon_data_extractor import cli
from amazon_data_extractor.asin_reader import InputWorkbookError, read_asins

ASIN = "B000000001"
OTHER_ASIN = "B000000002"
HEADERS = ["Status", "SKU", "Parentage Level", "Product Id Type", "Product Id", "Item Name"]
MACHINE = ["::listing_status", "contribution_sku#1.value", "parentage_level#1.value",
           "amzn1.volt.ca.product_id_type", "amzn1.volt.ca.product_id_value", "item_name#1.value"]


def report_bytes(rows=None, *, machine_only=False) -> bytes:
    workbook = Workbook()
    workbook.active.title = "Instructions"
    workbook.active.append(["Read the report instructions"])
    definitions = workbook.create_sheet("Data Definitions")
    definitions.append(["ASIN"])
    definitions.append(["B0DUMMY0001"])
    sheet = workbook.create_sheet("Template")
    if machine_only:
        sheet.append(MACHINE)
    else:
        sheet.append(["Category Listings Report"])
        sheet.append(["Instructions and version metadata"])
        sheet.append(["Grouped headings"])
        sheet.append(HEADERS)
        sheet.append(MACHINE)
        sheet.append([" ", " ", "parent", " ", " ", " "])
    for row in rows if rows is not None else [["Active", "SKU-1", "child", "ASIN", ASIN, "Synthetic product"]]:
        sheet.append(row)
    result = io.BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


def write_book(path, headers, rows, *, header_row=1, sheet_name="Products"):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = sheet_name
    for _ in range(header_row - 1):
        sheet.append(["Metadata"])
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    workbook.save(path)
    workbook.close()


class CategoryInputTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def test_simple_asin_workbooks_remain_supported_and_audited(self):
        path = self.root / "input.xlsx"
        write_book(path, ["ASIN", "SKU"], [[ASIN, "A"], [ASIN.lower(), "B"],
                   [f"https://www.amazon.com/dp/{OTHER_ASIN}", "C"], ["bad", "D"], ["", "E"]])
        result = read_asins(path)
        self.assertEqual(result.asins, [ASIN, OTHER_ASIN])
        self.assertEqual(result.header_row, 1)
        self.assertEqual(result.duplicate_count, 1)
        self.assertEqual(result.invalid_count, 2)

    def test_category_xlsx_and_xlsm_detect_template_and_header_row(self):
        for suffix in (".xlsx", ".xlsm"):
            with self.subTest(suffix=suffix):
                path = self.root / ("category" + suffix)
                path.write_bytes(report_bytes())
                result = read_asins(path)
                self.assertEqual(result.asins, [ASIN])
                self.assertEqual((result.sheet_name, result.header_row, result.data_start_row), ("Template", 4, 7))
                self.assertEqual(result.column_name, "Product Id")
                self.assertEqual(result.input_kind, "Category Listings Report")
                self.assertEqual(result.audit_rows[0].source_row, 7)

    def test_only_asin_typed_ids_are_submitted_and_parents_are_skipped(self):
        path = self.root / "category.xlsm"
        path.write_bytes(report_bytes([
            ["Active", "A", "child", "ASIN", ASIN, "One"],
            ["Active", "DUP", "child", "asin", ASIN.lower(), "Duplicate"],
            ["Active", "UPC-SKU", "child", "UPC", "0123456789", "Barcode"],
            ["Active", "PARENT-SKU", "parent", "ASIN", OTHER_ASIN, "Parent"],
            ["Active", "BLANK-SKU", "child", "ASIN", "", "Missing"],
            ["Active", "B", "child", "ASIN", OTHER_ASIN, "Two"],
        ]))
        result = read_asins(path)
        self.assertEqual(result.asins, [ASIN, OTHER_ASIN])
        self.assertEqual((result.duplicate_count, result.invalid_count, result.parent_count), (1, 2, 1))
        self.assertEqual(result.audit_rows[2].source_sku, "UPC-SKU")
        self.assertEqual(result.audit_rows[2].identifier_type, "UPC")
        self.assertIn("not ASIN", result.audit_rows[2].message)

    def test_technical_headers_work_without_human_headers(self):
        path = self.root / "technical.xlsx"
        path.write_bytes(report_bytes(machine_only=True))
        result = read_asins(path)
        self.assertEqual(result.asins, [ASIN])
        self.assertEqual(result.header_row, 1)

    def test_amazon_reserved_example_is_not_submitted_even_if_it_contains_an_asin(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Template"
        for _ in range(3):
            sheet.append(["Metadata"])
        sheet.append(HEADERS)
        sheet.append(MACHINE)
        sheet.append(["Active", "SAMPLE-SKU", "child", "ASIN", "B0EXAMPLE01", "Synthetic example"])
        sheet.append(["Active", "REAL-SKU", "child", "ASIN", ASIN, "Synthetic listing"])
        path = self.root / "with_example.xlsm"
        workbook.save(path)
        workbook.close()
        result = read_asins(path)
        self.assertEqual(result.asins, [ASIN])
        self.assertEqual(result.audit_rows[0].source_row, 7)

    def test_product_id_only_input_does_not_require_bullet_columns(self):
        path = self.root / "ids.xlsx"
        write_book(path, ["SKU", "Product ID Type", "Product ID"], [["A", "ASIN", ASIN]])
        self.assertEqual(read_asins(path).asins, [ASIN])

    def test_upc_ean_gtin_and_untyped_ids_are_not_guessed(self):
        path = self.root / "barcodes.xlsx"
        write_book(path, ["SKU", "Product ID Type", "Product ID"],
                   [[kind, kind, "0123456789"] for kind in ("UPC", "EAN", "GTIN", "GTIN Exempt", "")])
        result = read_asins(path)
        self.assertEqual(result.asins, [])
        self.assertEqual(result.invalid_count, 5)

    def test_manual_product_id_override_still_rejects_barcodes(self):
        path = self.root / "manual_ids.xlsx"
        write_book(path, ["SKU", "Product Id Type", "Product Id"],
                   [["A", "UPC", "0123456789"], ["B", "ASIN", ASIN]])
        result = read_asins(path, column_name="Product Id")
        self.assertEqual(result.asins, [ASIN])
        self.assertEqual(result.invalid_count, 1)

    def test_lowercase_template_name_also_skips_the_reserved_example(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "template"
        for _ in range(3):
            sheet.append(["Metadata"])
        sheet.append(HEADERS)
        sheet.append(MACHINE)
        sheet.append(["Active", "SAMPLE", "child", "ASIN", "B0EXAMPLE01", "Example"])
        sheet.append(["Active", "A", "child", "ASIN", ASIN, "Synthetic"])
        path = self.root / "lowercase.xlsx"
        workbook.save(path)
        workbook.close()
        self.assertEqual(read_asins(path).asins, [ASIN])

    def test_explicit_child_asin_wins_over_unrelated_barcode_and_parent_asin(self):
        path = self.root / "child.xlsx"
        write_book(path, ["Parent ASIN", "Child ASIN", "Product Id Type", "Product Id"],
                   [[OTHER_ASIN, ASIN, "UPC", "0123456789"]])
        self.assertEqual(read_asins(path).asins, [ASIN])

    def test_conflicting_asins_are_flagged_instead_of_silently_selected(self):
        path = self.root / "conflict.xlsx"
        write_book(path, ["ASIN", "Product Id Type", "Product Id"], [[ASIN, "ASIN", OTHER_ASIN]])
        result = read_asins(path)
        self.assertEqual(result.asins, [])
        self.assertIn("Conflicting", result.audit_rows[0].message)

    def test_manual_sheet_header_and_column_overrides_remain_supported(self):
        path = self.root / "custom.xlsx"
        write_book(path, ["Lookup"], [[ASIN]], header_row=4, sheet_name="Custom")
        self.assertEqual(read_asins(path, sheet_name="Custom", column_name="Lookup", header_row=4).asins, [ASIN])
        with self.assertRaises(InputWorkbookError):
            read_asins(path, header_row=0)

    def test_input_workbook_bytes_are_never_changed(self):
        path = self.root / "unchanged.xlsm"
        path.write_bytes(report_bytes())
        original = hashlib.sha256(path.read_bytes()).digest()
        read_asins(path)
        self.assertEqual(hashlib.sha256(path.read_bytes()).digest(), original)

    def test_helper_sheets_are_not_treated_as_listing_data(self):
        path = self.root / "helpers.xlsx"
        write_book(path, ["ASIN"], [[ASIN]], sheet_name="Data Definitions")
        with self.assertRaisesRegex(InputWorkbookError, "No ASIN column"):
            read_asins(path)

    def test_wrong_format_or_unrecognized_layout_has_a_clear_error(self):
        path = self.root / "legacy.xls"
        path.write_bytes(b"not a supported workbook")
        with self.assertRaisesRegex(InputWorkbookError, "xlsx or .xlsm"):
            read_asins(path)


class CategoryWorkflowTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        environment = patch.dict(os.environ, {"BRIGHTDATA_API_KEY": "synthetic-offline-key"})
        environment.start()
        self.addCleanup(environment.stop)
        self.client = create_app({"TESTING": True, "PROJECT_DIR": self.root}).test_client()

    def test_live_form_accepts_xlsm_and_audits_category_input(self):
        api = Mock()
        api.collect_sync.return_value = [{"asin": ASIN, "title": "Synthetic product"}]
        rows = [["Active", "SKU-A", "child", "ASIN", ASIN, "One"],
                ["Active", "SKU-UPC", "child", "UPC", "0123456789", "Barcode"]]
        with patch("app.BrightDataClient", return_value=api):
            response = self.client.post("/extract", data={
                "workbook": (io.BytesIO(report_bytes(rows)), "category.xlsm"), "mode": "Automatic",
            })
        self.assertEqual(response.status_code, 200)
        api.collect_sync.assert_called_once_with([ASIN])
        self.assertIn(b"SKU-UPC", response.data)
        self.assertIn(b"Rows without usable ASINs", response.data)
        metadata = json.loads(next(self.root.glob("output/*/run_metadata.json")).read_text())
        self.assertEqual(metadata["source_header_row"], 4)
        self.assertEqual(metadata["source_data_start_row"], 7)
        workbook = load_workbook(next(self.root.glob("output/*/amazon_extraction.xlsx")))
        try:
            audit = list(workbook["Input Audit"].values)
            self.assertEqual(audit[0][-3:], ("Source SKU", "Identifier Type", "Source Field"))
            self.assertEqual(audit[2][5:7], ("SKU-UPC", "UPC"))
        finally:
            workbook.close()

    def test_large_category_input_uses_async_and_unique_asins(self):
        asins = [f"B0{i:08d}" for i in range(25)]
        rows = [["Active", f"SKU-{i}", "child", "ASIN", asin, "Synthetic"] for i, asin in enumerate(asins)]
        api = Mock()
        api.trigger_collection.return_value = "s_offline"
        api.download_results.return_value = [{"asin": asin, "title": "Synthetic"} for asin in asins]
        with patch("app.BrightDataClient", return_value=api):
            response = self.client.post("/extract", data={
                "workbook": (io.BytesIO(report_bytes(rows)), "category.xlsm"), "mode": "Automatic",
            })
        self.assertEqual(response.status_code, 200)
        api.trigger_collection.assert_called_once_with(asins)
        api.collect_sync.assert_not_called()

    def test_no_usable_asins_blocks_the_provider_request(self):
        with patch("app.BrightDataClient") as provider:
            response = self.client.post("/extract", data={
                "workbook": (io.BytesIO(report_bytes([["Active", "A", "child", "EAN", "0123456789", "Barcode"]])), "category.xlsm"),
            }, follow_redirects=True)
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Bright Data was not called", response.data)
        provider.assert_not_called()

    def test_cli_dry_run_accepts_a_category_xlsm_without_api_calls(self):
        path = self.root / "category.xlsm"
        path.write_bytes(report_bytes())
        with patch.object(cli, "BrightDataClient") as provider, contextlib.redirect_stdout(io.StringIO()) as output:
            code = cli.main([str(path), "--dry-run"])
        self.assertEqual(code, 0)
        self.assertIn("Category Listings Report", output.getvalue())
        self.assertIn("Valid unique ASINs: 1", output.getvalue())
        provider.assert_not_called()


if __name__ == "__main__":
    unittest.main()
