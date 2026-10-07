from __future__ import annotations

from io import BytesIO
import json
import re
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from app import create_app


ASIN = "B000000001"


def workbook_bytes(build_workbook) -> BytesIO:
    with tempfile.NamedTemporaryFile(suffix=".xlsx") as temp_file:
        build_workbook.save(temp_file.name)
        return BytesIO(Path(temp_file.name).read_bytes())


def extraction_file() -> BytesIO:
    workbook = Workbook()
    data = workbook.active
    data.title = "Amazon Data"
    data.append(["input_asin", "asin", "title", "features", "url"])
    data.append(
        [
            ASIN,
            ASIN,
            "Example Product",
            json.dumps(["Exact bullet", "Live wording"]),
            f"https://www.amazon.com/dp/{ASIN}",
        ]
    )
    audit = workbook.create_sheet("Input Audit")
    audit.append(["Source Row", "Original Value", "Normalized ASIN", "Status", "Message"])
    audit.append([2, ASIN, ASIN, "ACCEPTED", ""])
    return workbook_bytes(workbook)


def catalog_file() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Template"
    sheet.append(
        ["SKU", "ASIN", "Title", "Bullet Point", "Bullet Point", "Bullet Point", "Bullet Point", "Bullet Point"]
    )
    sheet.append(["SKU-001", ASIN, "Example Product", "Exact bullet", "Catalog wording", "", "", ""])
    return workbook_bytes(workbook)


def catalog_missing_asin_file() -> BytesIO:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Catalog Data"
    sheet.append(
        [
            "SKU",
            "Product ID Type",
            "Product ID",
            "Product Name",
            "Bullet Point 1",
            "Bullet Point 2",
        ]
    )
    sheet.append(
        [
            "SKU-NEEDS-ASIN",
            "EAN",
            "00850072092171",
            "Catalog Product",
            "Exact bullet",
            "Catalog wording",
        ]
    )
    return workbook_bytes(workbook)


class AppTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.exit_calls = []
        self.app = create_app(
            {
                "TESTING": True,
                "PROJECT_DIR": Path(self.temp_dir.name),
                "SECRET_KEY": "test-key",
                "EXIT_SCHEDULER": lambda: self.exit_calls.append("scheduled"),
            }
        )
        self.client = self.app.test_client()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_index_contains_both_tabs(self):
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"Live Extraction", response.data)
        self.assertIn(b"Offline Compare", response.data)
        self.assertIn(b"All Listings Report", response.data)

    def test_exit_route_schedules_process_stop(self):
        response = self.client.post("/exit", environ_base={"REMOTE_ADDR": "127.0.0.1"})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(self.exit_calls, ["scheduled"])
        self.assertIn(b"Program closed", response.data)

    def test_offline_compare_creates_classified_report(self):
        response = self.client.post(
            "/compare",
            data={
                "extraction_workbook": (extraction_file(), "amazon_extraction.xlsx"),
                "catalog_workbook": (catalog_file(), "category_listings.xlsx"),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"OFFLINE COMPARISON COMPLETE", response.data)

        reports = list(Path(self.temp_dir.name).glob("output/amazon_comparison_*/comparison_report.xlsx"))
        self.assertEqual(len(reports), 1)
        workbook = load_workbook(reports[0], data_only=True)
        try:
            statuses = [
                row[7].value
                for row in workbook["Bullet Results"].iter_rows(min_row=2)
            ]
            self.assertEqual(statuses, ["Exact Match", "Text Changed"])
        finally:
            workbook.close()

    def test_missing_catalog_asin_asks_user_and_shows_identifier(self):
        response = self.client.post(
            "/compare",
            data={
                "extraction_workbook": (extraction_file(), "amazon_extraction.xlsx"),
                "catalog_workbook": (
                    catalog_missing_asin_file(),
                    "category_listings.xlsx",
                ),
            },
            content_type="multipart/form-data",
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"ASIN CONFIRMATION", response.data)
        self.assertIn(b"00850072092171", response.data)
        self.assertIn(b"SKU-NEEDS-ASIN", response.data)

        match = re.search(rb'action="/compare/([^/]+)/resolve"', response.data)
        self.assertIsNotNone(match)
        run_name = match.group(1).decode()
        response = self.client.post(
            f"/compare/{run_name}/resolve",
            data={
                "asin_2": ASIN,
                "extraction_name": "amazon_extraction.xlsx",
                "catalog_name": "category_listings.xlsx",
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"OFFLINE COMPARISON COMPLETE", response.data)

        report_path = (
            Path(self.temp_dir.name)
            / "output"
            / run_name
            / "comparison_report.xlsx"
        )
        workbook = load_workbook(report_path, data_only=True)
        try:
            details = workbook["Difference Details"]
            self.assertEqual(details["A1"].value, "ASIN")
            self.assertEqual(details["B1"].value, "Product Name")
            self.assertEqual(details["A2"].value, ASIN)
            self.assertEqual(details["B2"].value, "Catalog Product")
        finally:
            workbook.close()


if __name__ == "__main__":
    unittest.main()