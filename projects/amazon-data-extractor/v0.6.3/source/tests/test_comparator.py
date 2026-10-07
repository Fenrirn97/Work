from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from amazon_data_extractor.all_listings import AllListingsIndex
from amazon_data_extractor.catalog_reader import CatalogProduct, build_catalog_data
from amazon_data_extractor.comparator import (
    BULLET_MOVED,
    EXACT_MATCH,
    FORMATTING_ONLY,
    MISSING_CATALOG,
    MISSING_LIVE,
    TEXT_CHANGED,
    ComparisonResult,
    compare_bullets,
    resolve_catalog_asins,
)
from amazon_data_extractor.comparison_report import write_comparison_report
from amazon_data_extractor.live_records import LiveProduct


ASIN = "B000000001"


def live(*bullets: str) -> LiveProduct:
    values = (list(bullets) + [""] * 5)[:5]
    return LiveProduct(
        requested_asin=ASIN,
        returned_asin=ASIN,
        title="Example Product",
        bullets=values,
        page_url=f"https://www.amazon.com/dp/{ASIN}",
        status="Found",
        message="Found",
    )


def catalog(*bullets: str) -> CatalogProduct:
    return CatalogProduct(
        asin=ASIN,
        sku="SKU-001",
        product_name="Example Product",
        bullets=(list(bullets) + [""] * 5)[:5],
    )


class ComparatorTests(unittest.TestCase):
    def test_classifies_exact_formatting_changed_and_missing(self):
        results = compare_bullets(
            live("Exact", "fast charging", "Battery 4500", "", "Live only"),
            catalog("Exact", "FAST charging.", "Battery 5000", "Catalog only", ""),
        )
        self.assertEqual(
            [item.status for item in results],
            [EXACT_MATCH, FORMATTING_ONLY, TEXT_CHANGED, MISSING_LIVE, MISSING_CATALOG],
        )

    def test_detects_a_swap_as_two_moves_without_text_changes(self):
        results = compare_bullets(live("Second", "First"), catalog("First", "Second"))
        self.assertEqual([item.status for item in results], [BULLET_MOVED, BULLET_MOVED])
        self.assertEqual(
            {(item.catalog_position, item.live_position) for item in results},
            {(1, 2), (2, 1)},
        )

    def test_detects_inserted_live_bullet_without_double_counting(self):
        results = compare_bullets(live("First", "Inserted", "Second"), catalog("First", "Second"))
        self.assertEqual(
            [item.status for item in results],
            [EXACT_MATCH, MISSING_CATALOG, BULLET_MOVED],
        )

    def test_report_contains_all_and_review_only_sheets(self):
        comparisons = compare_bullets(
            live("Exact", "Different live"),
            catalog("Exact", "Different catalog"),
        )
        result = ComparisonResult(
            comparisons=comparisons,
            live_problems=[],
            not_in_catalog=[],
            stats={},
        )
        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = Path(temp_dir) / "comparison_report.xlsx"
            write_comparison_report(report_path, result)
            workbook = load_workbook(report_path, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "ASIN Issues",
                        "Bullet Results",
                        "Difference Details",
                        "ASIN Needed",
                        "Not Found in Catalog",
                        "Live Extraction Problems",
                        "Read Me",
                    ],
                )
                all_statuses = [
                    row[7].value
                    for row in workbook["Bullet Results"].iter_rows(min_row=2)
                ]
                issue_statuses = [
                    row[7].value
                    for row in workbook["Difference Details"].iter_rows(min_row=2)
                ]
                self.assertEqual(all_statuses, [EXACT_MATCH, TEXT_CHANGED])
                self.assertEqual(issue_statuses, [TEXT_CHANGED])
                details = workbook["Difference Details"]
                self.assertEqual(details["A1"].value, "ASIN")
                self.assertEqual(details["B1"].value, "Product Name")
                self.assertEqual(details["A2"].value, ASIN)
                self.assertEqual(details["B2"].value, "Example Product")
            finally:
                workbook.close()

    def test_resolves_blank_asins_by_sku_name_then_user(self):
        products = [
            CatalogProduct(
                "",
                "SKU-BRIDGE",
                "Different Name",
                [""] * 5,
                "EAN",
                "1000000000001",
                "",
                7,
            ),
            CatalogProduct(
                "",
                "SKU-NAME",
                "Example Product",
                [""] * 5,
                "GTIN",
                "1000000000002",
                "",
                8,
            ),
            CatalogProduct(
                "",
                "SKU-USER",
                "Ambiguous Name",
                [""] * 5,
                "EAN",
                "1000000000003",
                "",
                9,
            ),
        ]
        catalog_data = build_catalog_data(products, sheet_name="Template")
        listings = AllListingsIndex(
            asin_to_skus={ASIN: ["SKU-BRIDGE"]},
            sku_to_asin={"sku-bridge": ASIN},
            row_count=1,
        )

        resolved, stats, unresolved = resolve_catalog_asins(
            catalog_data,
            {ASIN: live("One")},
            listings,
            {9: "B000000009"},
        )

        self.assertEqual(
            [product.asin for product in resolved.products],
            [ASIN, ASIN, "B000000009"],
        )
        self.assertEqual(stats["catalog_asins_resolved_by_sku"], 1)
        self.assertEqual(stats["catalog_asins_resolved_by_name"], 1)
        self.assertEqual(stats["catalog_asins_resolved_by_user"], 1)
        self.assertEqual(unresolved, [])

    def test_asin_needed_preserves_identifier_number(self):
        unresolved = CatalogProduct(
            "",
            "SKU-NEEDS-ASIN",
            "Needs ASIN",
            [""] * 5,
            "EAN",
            "00850072092171",
            "",
            17,
        )
        result = ComparisonResult([], [], [], {}, [unresolved])
        with tempfile.TemporaryDirectory() as temp_dir:
            report_path = Path(temp_dir) / "comparison_report.xlsx"
            write_comparison_report(report_path, result)
            workbook = load_workbook(report_path, data_only=True)
            try:
                row = next(
                    workbook["ASIN Needed"].iter_rows(
                        min_row=2, max_row=2, values_only=True
                    )
                )
                self.assertEqual(
                    row[:4],
                    ("EAN", "00850072092171", "SKU-NEEDS-ASIN", "Needs ASIN"),
                )
            finally:
                workbook.close()


if __name__ == "__main__":
    unittest.main()