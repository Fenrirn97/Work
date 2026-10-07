"""Write the v0.6.3 bullet-classification workbook."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")
STATUS_FILLS = {
    "Exact Match": PatternFill("solid", fgColor="E2F0D9"),
    "Formatting Only": PatternFill("solid", fgColor="DDEBF7"),
    "Text Changed": PatternFill("solid", fgColor="F4CCCC"),
    "Missing Live": PatternFill("solid", fgColor="FCE4D6"),
    "Missing Catalog": PatternFill("solid", fgColor="FCE4D6"),
    "Bullet Moved": PatternFill("solid", fgColor="FCE5CD"),
}
STATUS_ORDER = {
    "Text Changed": 0,
    "Missing Live": 1,
    "Missing Catalog": 2,
    "Bullet Moved": 3,
    "Formatting Only": 4,
    "Exact Match": 5,
}


def amazon_url(asin: str) -> str:
    return f"https://www.amazon.com/dp/{asin}" if asin else ""


def add_amazon_link(cell, asin: str) -> None:
    cell.value = "Open Amazon"
    cell.hyperlink = amazon_url(asin)
    cell.style = "Hyperlink"


def format_sheet(worksheet, widths: list[int]) -> None:
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 30

    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in worksheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    for number, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(number)].width = width


def action_for(status: str) -> str:
    return {
        "Exact Match": "No action",
        "Formatting Only": "Content matches; formatting review is optional",
        "Text Changed": "Review the wording",
        "Missing Live": "Add or restore the live bullet",
        "Missing Catalog": "Review whether the catalog needs this bullet",
        "Bullet Moved": "Review bullet order",
    }[status]


def excel_safe(value):
    """Keep uploaded text from becoming an Excel formula."""
    if not isinstance(value, str):
        return value
    if value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def append_comparison_row(worksheet, item) -> None:
    worksheet.append(
        [
            item.asin,
            excel_safe(item.product_name),
            excel_safe(item.sku),
            excel_safe(item.matched_via),
            "",
            item.catalog_field,
            item.live_field,
            item.status,
            excel_safe(item.catalog_value),
            excel_safe(item.live_value),
            "Yes" if item.needs_review else "No",
            action_for(item.status),
        ]
    )
    row_number = worksheet.max_row
    add_amazon_link(worksheet.cell(row_number, 5), item.asin)
    worksheet.cell(row_number, 8).fill = STATUS_FILLS[item.status]
    if item.needs_review:
        worksheet.cell(row_number, 11).fill = REVIEW_FILL


def write_comparison_report(output_path, comparison) -> None:
    workbook = Workbook()
    issues_sheet = workbook.active
    issues_sheet.title = "ASIN Issues"
    results_sheet = workbook.create_sheet("Bullet Results")
    details_sheet = workbook.create_sheet("Difference Details")
    asin_needed_sheet = workbook.create_sheet("ASIN Needed")
    not_catalog_sheet = workbook.create_sheet("Not Found in Catalog")
    live_problem_sheet = workbook.create_sheet("Live Extraction Problems")
    guide_sheet = workbook.create_sheet("Read Me")

    issues_sheet.append(
        [
            "ASIN",
            "Product Name",
            "SKU(s)",
            "Amazon Link",
            "Issue Count",
            "Classifications",
            "Status",
        ]
    )
    for asin, issues in comparison.issues_by_asin.items():
        skus = sorted({issue.sku for issue in issues if issue.sku}, key=str.casefold)
        name = next((issue.product_name for issue in issues if issue.product_name), "")
        statuses = sorted(
            {issue.status for issue in issues},
            key=lambda status: STATUS_ORDER.get(status, 99),
        )
        issues_sheet.append(
            [
                asin,
                excel_safe(name),
                excel_safe(", ".join(skus)),
                "",
                len(issues),
                ", ".join(statuses),
                "Needs Review",
            ]
        )
        add_amazon_link(issues_sheet.cell(issues_sheet.max_row, 4), asin)
        issues_sheet.cell(issues_sheet.max_row, 7).fill = REVIEW_FILL

    result_headers = [
        "ASIN",
        "Product Name",
        "SKU",
        "Matched Via",
        "Amazon Link",
        "Catalog Field",
        "Live Field",
        "Classification",
        "Catalog Value",
        "Live Value",
        "Needs Review",
        "Action",
    ]
    results_sheet.append(result_headers)
    for item in comparison.comparisons:
        append_comparison_row(results_sheet, item)

    details_sheet.append(result_headers)
    for item in comparison.issues:
        append_comparison_row(details_sheet, item)

    asin_needed_sheet.append(
        [
            "Identifier Type",
            "Identifier Number",
            "SKU",
            "Product Name",
            "Catalog Row",
            "Status",
        ]
    )
    for product in comparison.unresolved_catalog:
        asin_needed_sheet.append(
            [
                excel_safe(product.identifier_type or "Unknown"),
                excel_safe(product.identifier_value or "(blank)"),
                excel_safe(product.sku),
                excel_safe(product.product_name),
                product.source_row,
                "ASIN not provided",
            ]
        )

    not_catalog_sheet.append(["ASIN", "Live Product Name", "Amazon Link", "Status"])
    for product in comparison.not_in_catalog:
        not_catalog_sheet.append(
            [
                product.requested_asin,
                excel_safe(product.title),
                "",
                "Live ASIN was found, but it was not found in the Category Listings Report.",
            ]
        )
        add_amazon_link(
            not_catalog_sheet.cell(not_catalog_sheet.max_row, 3),
            product.requested_asin,
        )
        not_catalog_sheet.cell(not_catalog_sheet.max_row, 4).fill = REVIEW_FILL

    live_problem_sheet.append(
        ["Requested ASIN", "Returned ASIN", "Amazon Link", "Result", "What Happened", "Action"]
    )
    for product in comparison.live_problems:
        live_problem_sheet.append(
            [
                product.requested_asin,
                product.returned_asin,
                "",
                product.status,
                excel_safe(product.message),
                "Review before comparing" if product.status == "Redirected" else "Retry or review",
            ]
        )
        add_amazon_link(
            live_problem_sheet.cell(live_problem_sheet.max_row, 3),
            product.requested_asin,
        )
        live_problem_sheet.cell(live_problem_sheet.max_row, 4).fill = ERROR_FILL

    guide_sheet.append(["Classification", "Meaning", "Needs Review"])
    guide_rows = [
        ("Exact Match", "The text is identical in the same bullet position.", "No"),
        (
            "Formatting Only",
            "The normalized content matches; only capitalization, punctuation, "
            "spacing, HTML, or similar presentation differs.",
            "No",
        ),
        ("Text Changed", "Both positions contain text, but the content is different.", "Yes"),
        ("Missing Live", "The catalog contains a bullet that the live result does not.", "Yes"),
        ("Missing Catalog", "The live result contains a bullet that the catalog does not.", "Yes"),
        ("Bullet Moved", "The same normalized bullet exists in both files at different positions.", "Yes"),
    ]
    for row in guide_rows:
        guide_sheet.append(row)
        guide_sheet.cell(guide_sheet.max_row, 1).fill = STATUS_FILLS[row[0]]

    guide_sheet.append([])
    guide_sheet.append(
        [
            "Product matching order",
            "ASIN; returned ASIN; SKU through the optional All Listings TXT; "
            "unique product name; user-entered ASIN.",
            "",
        ]
    )
    guide_sheet.append(
        [
            "ASIN Needed",
            "Rows left blank on the ASIN question screen remain here with their "
            "original identifier number.",
            "",
        ]
    )

    format_sheet(issues_sheet, [16, 62, 38, 18, 12, 34, 18])
    format_sheet(results_sheet, [16, 48, 32, 28, 18, 18, 18, 20, 72, 72, 15, 42])
    format_sheet(details_sheet, [16, 48, 32, 28, 18, 18, 18, 20, 72, 72, 15, 42])
    format_sheet(asin_needed_sheet, [20, 26, 42, 62, 14, 24])
    format_sheet(not_catalog_sheet, [16, 62, 18, 72])
    format_sheet(live_problem_sheet, [16, 16, 18, 20, 72, 28])
    format_sheet(guide_sheet, [22, 96, 18])

    for worksheet in (
        issues_sheet,
        results_sheet,
        details_sheet,
        asin_needed_sheet,
        not_catalog_sheet,
        live_problem_sheet,
    ):
        for row in worksheet.iter_rows(min_row=2):
            row[0].number_format = "@"
            if worksheet is asin_needed_sheet and len(row) > 1:
                row[1].number_format = "@"

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)