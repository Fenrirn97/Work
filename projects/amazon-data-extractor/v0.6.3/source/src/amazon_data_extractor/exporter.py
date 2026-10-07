"""Write the untouched JSON response and a human-readable Excel workbook."""


from __future__ import annotations


import json
from dataclasses import asdict
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable


from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


from .asin_reader import AsinAuditRow




# Stable field order from the confirmed Bright Data Amazon product schema.
# Unknown fields are appended, so schema additions are not discarded.
KNOWN_FIELDS = (
    "category_tree", "target_countries", "category_urls", "title", "seller_name",
    "brand", "description", "initial_price", "currency", "availability",
    "reviews_count", "categories", "parent_asin", "asin", "buybox_seller",
    "number_of_sellers", "root_bs_rank", "ISBN10", "answered_questions", "domain",
    "images_count", "url", "video_count", "image_url", "item_weight", "rating",
    "product_dimensions", "seller_id", "image", "date_first_available", "discount",
    "model_number", "manufacturer", "department", "plus_content", "upc", "video",
    "top_review", "final_price_high", "final_price", "variations", "delivery",
    "features", "format", "buybox_prices", "input_asin", "ingredients", "origin_url",
    "bought_past_month", "is_available", "root_bs_category", "bs_category", "bs_rank",
    "badge", "subcategory_rank", "amazon_choice", "images", "product_details",
    "prices_breakdown", "country_of_origin", "from_the_brand", "product_description",
    "seller_url", "customer_says", "sustainability_features", "climate_pledge_friendly",
    "videos", "other_sellers_prices", "downloadable_videos", "editorial_reviews",
    "about_the_author", "zipcode", "coupon", "sponsered", "store_url", "ships_from",
    "city", "customers_say", "max_quantity_available", "variations_values", "language",
    "return_policy", "inactive_buy_box", "buybox_seller_rating", "premium_brand",
    "amazon_prime", "coupon_description", "all_badges", "sponsored",
    "variant_attributes", "safety_information", "subcategory_link",
    "is_frequently_returned_item_badge", "frequently_returned_item_message",
    "is_customers_usually_keep", "title_badge", "review_images", "review_videos",
    "bought_past_month_text", "is_high_price", "title_highlight", "title_clean",
    "customers_say_topics", "timestamp", "input_url", "input_language",
    "input_all_variations", "input_zipcode", "error", "error_code", "warning",
    "warning_code",
)


WORKBOOK_SCHEMA_VERSION = "v2.0"
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
ALTERNATE_FILL = PatternFill("solid", fgColor="F3F7FB")
LINK_FONT = Font(color="0563C1", underline="single")
HEADER_BORDER = Border(bottom=Side(style="medium", color="17365D"))
EXCEL_CELL_LIMIT = 32_767




PRODUCT_HEADERS = (
    "Requested ASIN", "ASIN", "Parent ASIN", "Complete Amazon Title", "Brand",
    "Model Number", "Current Price", "Original Price", "Currency", "Availability",
    "Available", "Maximum Orderable Quantity", "Seller", "Ships From",
    "Buy Box Seller", "Number of Sellers", "Prime Eligible", "Rating", "Reviews",
    "Bought in Past Month", "Main Category", "Main Category Rank", "Amazon's Choice",
    "Product Dimensions", "Item Weight", "Country of Origin", "Main Image",
    "Amazon Page", "Collected At",
)




def _json_default(value: object) -> str:
    if isinstance(value, datetime):
        return value.isoformat()
    return str(value)




def _excel_value(value: Any) -> str | int | float | bool:
    """Convert a value safely without allowing spreadsheet formula injection."""


    if value is None:
        return ""
    if isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, (dict, list, tuple)):
        text = json.dumps(value, ensure_ascii=False, default=_json_default)
    else:
        text = str(value)

    if len(text) > EXCEL_CELL_LIMIT:
        text = text[:32_680] + "... [complete value is in brightdata_raw.json]"
    if text.startswith(("=", "+", "-", "@")):
        text = "'" + text
    return text




def _decode_nested(value: Any) -> Any:
    """Decode JSON-looking strings while leaving ordinary text unchanged."""


    if not isinstance(value, str):
        return value
    stripped = value.strip()
    if len(stripped) < 2 or stripped[0] not in "[{" or stripped[-1] not in "]}":
        return value
    try:
        return json.loads(stripped)
    except (TypeError, ValueError, json.JSONDecodeError):
        return value




def _items(value: Any) -> list[Any]:
    value = _decode_nested(value)
    if value in (None, ""):
        return []
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]




def _human_value(value: Any) -> str | int | float:
    """Make common nested values readable on the human-facing worksheets."""


    value = _decode_nested(value)
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, (list, tuple)):
        if all(not isinstance(item, (dict, list, tuple)) for item in value):
            return "\n".join(str(item) for item in value if item not in (None, ""))
        return _excel_value(value)
    if isinstance(value, dict):
        if all(not isinstance(item, (dict, list, tuple)) for item in value.values()):
            return "\n".join(
                f"{key}: {item}" for key, item in value.items() if item not in (None, "")
            )
        return _excel_value(value)
    return _excel_value(value)




def _ordered_fields(records: Iterable[dict[str, Any]]) -> list[str]:
    discovered: set[str] = set()
    for record in records:
        discovered.update(record)
    known = list(KNOWN_FIELDS)
    unknown = sorted(discovered.difference(KNOWN_FIELDS), key=str.lower)
    return known + unknown




def _first_value(record: dict[str, Any], *fields: str) -> Any:
    for field in fields:
        value = record.get(field)
        if value not in (None, ""):
            return value
    return ""




def _record_asin(record: dict[str, Any]) -> str:
    asin = _first_value(record, "asin", "input_asin")
    if asin:
        return str(asin)
    input_value = _decode_nested(record.get("input"))
    if isinstance(input_value, dict) and input_value.get("asin"):
        return str(input_value["asin"])
    return ""




def _requested_asin(record: dict[str, Any]) -> str:
    requested = record.get("input_asin")
    if requested:
        return str(requested)
    input_value = _decode_nested(record.get("input"))
    if isinstance(input_value, dict) and input_value.get("asin"):
        return str(input_value["asin"])
    return _record_asin(record)




def _category_name(record: dict[str, Any]) -> Any:
    root = record.get("root_bs_category")
    if root not in (None, ""):
        return root
    categories = _items(record.get("categories"))
    return categories[-1] if categories else ""




def _product_row(record: dict[str, Any]) -> list[Any]:
    return [
        _requested_asin(record), _record_asin(record), record.get("parent_asin"),
        record.get("title"), record.get("brand"), record.get("model_number"),
        _first_value(record, "final_price", "initial_price"),
        record.get("initial_price"), record.get("currency"),
        record.get("availability"), _human_value(record.get("is_available")),
        record.get("max_quantity_available"), record.get("seller_name"),
        record.get("ships_from"), record.get("buybox_seller"),
        record.get("number_of_sellers"), _human_value(record.get("amazon_prime")),
        record.get("rating"), record.get("reviews_count"),
        _first_value(record, "bought_past_month_text", "bought_past_month"),
        _category_name(record), record.get("root_bs_rank"),
        _human_value(record.get("amazon_choice")), record.get("product_dimensions"),
        record.get("item_weight"), record.get("country_of_origin"),
        _first_value(record, "image_url", "image"),
        _first_value(record, "url", "origin_url", "input_url"),
        record.get("timestamp"),
    ]




def _style_table(
    worksheet,
    *,
    freeze: str = "A2",
    max_width: int = 36,
    widths: dict[str, int] | None = None,
    wrap_headers: set[str] | None = None,
    hyperlink_headers: set[str] | None = None,
    number_formats: dict[str, str] | None = None,
) -> None:
    """Apply a consistent, readable style to a worksheet table."""


    widths = widths or {}
    wrap_headers = wrap_headers or set()
    hyperlink_headers = hyperlink_headers or set()
    number_formats = number_formats or {}
    worksheet.freeze_panes = freeze
    worksheet.auto_filter.ref = worksheet.dimensions
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 32


    headers = [str(cell.value or "") for cell in worksheet[1]]
    for cell in worksheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = HEADER_BORDER
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )


    for row_index in range(2, worksheet.max_row + 1):
        if row_index % 2 == 0:
            for cell in worksheet[row_index]:
                cell.fill = ALTERNATE_FILL
        for column_index, cell in enumerate(worksheet[row_index], start=1):
            header = headers[column_index - 1]
            cell.alignment = Alignment(vertical="top", wrap_text=header in wrap_headers)
            if header in number_formats and cell.value not in (None, ""):
                cell.number_format = number_formats[header]
            if (
                header in hyperlink_headers
                and isinstance(cell.value, str)
                and cell.value.startswith(("http://", "https://"))
            ):
                cell.hyperlink = cell.value
                cell.font = LINK_FONT


    for column_index, column_cells in enumerate(worksheet.iter_cols(), start=1):
        header = headers[column_index - 1]
        if header in widths:
            width = widths[header]
        else:
            longest = max(
                (
                    max((len(line) for line in str(cell.value or "").splitlines()), default=0)
                    for cell in column_cells[:100]
                ),
                default=8,
            )
            width = min(longest + 2, max_width)
        worksheet.column_dimensions[get_column_letter(column_index)].width = max(10, width)




def _use_short_link_labels(worksheet, labels: dict[str, str]) -> None:
    """Keep links clickable without filling the overview with long URLs."""


    headers = [str(cell.value or "") for cell in worksheet[1]]
    for header, label in labels.items():
        if header not in headers:
            continue
        column = headers.index(header) + 1
        for row_index in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row_index, column=column)
            if not isinstance(cell.value, str) or not cell.value.startswith(
                ("http://", "https://")
            ):
                continue
            url = cell.value
            cell.hyperlink = url
            cell.value = label
            cell.font = LINK_FONT




def _append_attribute_rows(
    rows: list[list[Any]], asin: str, value: Any, source: str
) -> None:
    value = _decode_nested(value)
    if value in (None, ""):
        return
    if isinstance(value, (list, tuple)):
        for item in value:
            _append_attribute_rows(rows, asin, item, source)
        return
    if isinstance(value, dict):
        label = _first_value(value, "type", "name", "label", "attribute", "key")
        if label:
            detail_value = _first_value(value, "value", "values", "title", "text")
            if detail_value == "":
                detail_value = {
                    key: item for key, item in value.items() if key not in {
                        "type", "name", "label", "attribute", "key"
                    }
                }
            rows.append([asin, label, _human_value(detail_value), source])
            return
        for key, item in value.items():
            rows.append([asin, key, _human_value(item), source])
        return
    rows.append([asin, source, _human_value(value), source])




def _image_url(value: Any) -> str:
    value = _decode_nested(value)
    if isinstance(value, str):
        return value
    if isinstance(value, dict):
        for key in ("url", "image_url", "src", "link"):
            if value.get(key):
                return str(value[key])
    return ""




def _variation_rows(record: dict[str, Any]) -> list[list[Any]]:
    rows: list[list[Any]] = []
    asin = _record_asin(record)
    number = 0
    for source in ("variations", "variations_values", "variant_attributes"):
        parsed = _decode_nested(record.get(source))
        if parsed in (None, ""):
            continue
        items = parsed if isinstance(parsed, (list, tuple)) else [parsed]
        for item in items:
            if isinstance(item, dict) and not any(
                key in item
                for key in (
                    "asin", "variation_asin", "child_asin", "name", "type",
                    "attribute", "value", "title", "label",
                )
            ):
                for attribute, value in item.items():
                    number += 1
                    rows.append(
                        [asin, source, number, "", attribute, _human_value(value), ""]
                    )
                continue


            number += 1
            if isinstance(item, dict):
                variation_asin = _first_value(
                    item, "asin", "variation_asin", "child_asin"
                )
                attribute = _first_value(item, "name", "type", "attribute")
                value = _first_value(item, "value", "title", "label", "values")
                details = _excel_value(item)
            else:
                variation_asin = ""
                attribute = ""
                value = _human_value(item)
                details = ""
            rows.append(
                [asin, source, number, variation_asin, attribute, value, details]
            )
    return rows




def write_outputs(
    run_directory: str | Path,
    *,
    records: list[dict[str, Any]],
    audit_rows: list[AsinAuditRow],
    metadata: dict[str, Any],
) -> tuple[Path, Path, Path]:
    """Write raw JSON, metadata, and the organized workbook."""


    run_dir = Path(run_directory)
    run_dir.mkdir(parents=True, exist_ok=True)
    raw_json_path = run_dir / "brightdata_raw.json"
    metadata_path = run_dir / "run_metadata.json"
    workbook_path = run_dir / "amazon_extraction.xlsx"
    complete_metadata = {
        "workbook_schema_version": WORKBOOK_SCHEMA_VERSION,
        "returned_record_count": len(records),
        **metadata,
    }


    raw_json_path.write_text(
        json.dumps(records, indent=2, ensure_ascii=False, default=_json_default),
        encoding="utf-8",
    )
    metadata_path.write_text(
        json.dumps(complete_metadata, indent=2, ensure_ascii=False, default=_json_default),
        encoding="utf-8",
    )


    workbook = Workbook()
    workbook.properties.title = "Amazon Data Extraction"
    workbook.properties.subject = "Bright Data Amazon product extraction"


    overview_sheet = workbook.active
    overview_sheet.title = "Product Overview"
    overview_sheet.append(PRODUCT_HEADERS)
    for record in records:
        overview_sheet.append([_excel_value(value) for value in _product_row(record)])
    _style_table(
        overview_sheet,
        max_width=28,
        widths={
            "Requested ASIN": 16, "ASIN": 16, "Parent ASIN": 16,
            "Complete Amazon Title": 55, "Availability": 20,
            "Maximum Orderable Quantity": 25, "Bought in Past Month": 24,
            "Product Dimensions": 24, "Main Image": 18, "Amazon Page": 20,
            "Collected At": 25,
        },
        wrap_headers={
            "Complete Amazon Title", "Availability", "Bought in Past Month",
            "Product Dimensions",
        },
        hyperlink_headers={"Main Image", "Amazon Page"},
        number_formats={
            "Current Price": "#,##0.00", "Original Price": "#,##0.00",
            "Maximum Orderable Quantity": "#,##0", "Number of Sellers": "#,##0",
            "Rating": "0.0", "Reviews": "#,##0", "Main Category Rank": "#,##0",
        },
    )
    _use_short_link_labels(
        overview_sheet,
        {"Main Image": "Open image", "Amazon Page": "Open Amazon page"},
    )
    for row_index in range(2, overview_sheet.max_row + 1):
        overview_sheet.row_dimensions[row_index].height = 45


    bullets_sheet = workbook.create_sheet("Bullets")
    bullets_sheet.append(["ASIN", "Bullet Number", "Bullet Point"])
    for record in records:
        asin = _record_asin(record)
        for index, bullet in enumerate(_items(record.get("features")), start=1):
            if isinstance(bullet, dict):
                bullet = _first_value(bullet, "text", "value", "title") or bullet
            bullets_sheet.append([asin, index, _human_value(bullet)])
    _style_table(
        bullets_sheet,
        widths={"ASIN": 16, "Bullet Number": 15, "Bullet Point": 110},
        wrap_headers={"Bullet Point"},
        number_formats={"Bullet Number": "0"},
    )


    details_sheet = workbook.create_sheet("Product Details")
    details_sheet.append(["ASIN", "Attribute", "Value", "Source"])
    detail_rows: list[list[Any]] = []
    for record in records:
        asin = _record_asin(record)
        _append_attribute_rows(
            detail_rows, asin, record.get("product_details"), "Product Details"
        )
        _append_attribute_rows(
            detail_rows, asin, record.get("variant_attributes"), "Variant Attributes"
        )
    unique_detail_rows: list[list[Any]] = []
    seen_details: set[tuple[str, ...]] = set()
    for row in detail_rows:
        identity = tuple(str(value) for value in row)
        if identity in seen_details:
            continue
        seen_details.add(identity)
        unique_detail_rows.append(row)
    for row in unique_detail_rows:
        details_sheet.append([_excel_value(value) for value in row])
    _style_table(
        details_sheet,
        widths={"ASIN": 16, "Attribute": 36, "Value": 70, "Source": 22},
        wrap_headers={"Value"},
    )


    images_sheet = workbook.create_sheet("Images")
    images_sheet.append(["ASIN", "Image Number", "Image Type", "Image URL"])
    for record in records:
        asin = _record_asin(record)
        seen: set[str] = set()
        image_number = 0
        image_groups = (
            ("Main", [_first_value(record, "image_url", "image")]),
            ("Gallery", _items(record.get("images"))),
            ("Review", _items(record.get("review_images"))),
        )
        for image_type, image_values in image_groups:
            for image_value in image_values:
                url = _image_url(image_value)
                if not url or url in seen:
                    continue
                seen.add(url)
                image_number += 1
                images_sheet.append([asin, image_number, image_type, url])
    _style_table(
        images_sheet,
        widths={"ASIN": 16, "Image Number": 15, "Image Type": 15, "Image URL": 100},
        hyperlink_headers={"Image URL"},
        number_formats={"Image Number": "0"},
    )


    categories_sheet = workbook.create_sheet("Categories")
    categories_sheet.append(
        ["ASIN", "Category Type", "Level", "Category", "Category URL", "Rank"]
    )
    for record in records:
        asin = _record_asin(record)
        category_tree = _items(record.get("category_tree"))
        if category_tree:
            for level, item in enumerate(category_tree, start=1):
                if isinstance(item, dict):
                    name = _first_value(item, "name", "category", "title")
                    url = _first_value(item, "url", "link")
                else:
                    name = item
                    url = ""
                categories_sheet.append([asin, "Browse Path", level, name, url, ""])
        else:
            names = _items(record.get("categories"))
            urls = _items(record.get("category_urls"))
            for level, name in enumerate(names, start=1):
                url = urls[level - 1] if level <= len(urls) else ""
                categories_sheet.append([asin, "Browse Path", level, name, url, ""])


        if record.get("root_bs_category") or record.get("root_bs_rank"):
            categories_sheet.append(
                [
                    asin, "Best Seller Rank", 1, record.get("root_bs_category"), "",
                    record.get("root_bs_rank"),
                ]
            )
        if record.get("bs_category") or record.get("bs_rank"):
            categories_sheet.append(
                [
                    asin, "Best Seller Rank", 2, record.get("bs_category"), "",
                    record.get("bs_rank"),
                ]
            )
        recorded_ranks = {
            (str(record.get("root_bs_category") or ""), str(record.get("root_bs_rank") or "")),
            (str(record.get("bs_category") or ""), str(record.get("bs_rank") or "")),
        }
        for item in _items(record.get("subcategory_rank")):
            if isinstance(item, dict):
                subcategory_name = _first_value(
                    item, "subcategory_name", "name", "category"
                )
                subcategory_rank = _first_value(item, "subcategory_rank", "rank")
                rank_identity = (str(subcategory_name), str(subcategory_rank))
                if rank_identity in recorded_ranks:
                    continue
                recorded_ranks.add(rank_identity)
                categories_sheet.append(
                    [
                        asin, "Best Seller Rank", "", subcategory_name,
                        _first_value(item, "subcategory_link", "url", "link"),
                        subcategory_rank,
                    ]
                )
    _style_table(
        categories_sheet,
        widths={
            "ASIN": 16, "Category Type": 20, "Level": 10, "Category": 38,
            "Category URL": 80, "Rank": 14,
        },
        hyperlink_headers={"Category URL"},
        number_formats={"Level": "0", "Rank": "#,##0"},
    )


    variations_sheet = workbook.create_sheet("Variations")
    variations_sheet.append(
        [
            "ASIN", "Source", "Variation Number", "Variation ASIN", "Attribute",
            "Value", "Complete Details",
        ]
    )
    for record in records:
        for row in _variation_rows(record):
            variations_sheet.append([_excel_value(value) for value in row])
    _style_table(
        variations_sheet,
        widths={
            "ASIN": 16, "Source": 22, "Variation Number": 18,
            "Variation ASIN": 18, "Attribute": 28, "Value": 42,
            "Complete Details": 80,
        },
        wrap_headers={"Value", "Complete Details"},
        number_formats={"Variation Number": "0"},
    )


    sellers_sheet = workbook.create_sheet("Sellers")
    sellers_sheet.append(
        [
            "ASIN", "Offer Type", "Seller", "Seller ID", "Price", "Currency",
            "Ships From", "Seller Rating", "Number of Ratings", "Delivery", "Seller URL",
        ]
    )
    for record in records:
        asin = _record_asin(record)
        main_seller = _first_value(record, "seller_name", "buybox_seller")
        main_price = _first_value(record, "final_price", "initial_price")
        if any(
            record.get(field) not in (None, "")
            for field in (
                "seller_name", "seller_id", "buybox_seller", "final_price",
                "ships_from", "seller_url",
            )
        ):
            offer_type = (
                "Buy Box"
                if record.get("buybox_seller")
                and record.get("buybox_seller") == record.get("seller_name")
                else "Current Offer"
            )
            sellers_sheet.append(
                [
                    asin, offer_type, main_seller, record.get("seller_id"), main_price,
                    record.get("currency"), record.get("ships_from"),
                    record.get("buybox_seller_rating"), "",
                    _human_value(record.get("delivery")), record.get("seller_url"),
                ]
            )
        for seller in _items(record.get("other_sellers_prices")):
            if not isinstance(seller, dict):
                continue
            seller_name = _first_value(seller, "seller_name", "seller")
            seller_price = _first_value(seller, "price", "final_price")
            if str(seller_name) == str(main_seller) and seller_price == main_price:
                continue
            sellers_sheet.append(
                [
                    asin, "Other Seller", seller_name,
                    _first_value(seller, "seller_id", "id"),
                    seller_price,
                    _first_value(seller, "currency") or record.get("currency"),
                    _first_value(seller, "ships_from", "shipper"),
                    _first_value(seller, "seller_rating", "rating"),
                    _first_value(seller, "num_of_ratings", "ratings_count"),
                    _human_value(seller.get("delivery")),
                    _first_value(seller, "seller_url", "url"),
                ]
            )
    _style_table(
        sellers_sheet,
        widths={
            "ASIN": 16, "Offer Type": 18, "Seller": 28, "Seller ID": 22,
            "Price": 14, "Currency": 12, "Ships From": 20, "Seller Rating": 16,
            "Number of Ratings": 18, "Delivery": 60, "Seller URL": 80,
        },
        wrap_headers={"Delivery"},
        hyperlink_headers={"Seller URL"},
        number_formats={
            "Price": "#,##0.00", "Seller Rating": "0.0",
            "Number of Ratings": "#,##0",
        },
    )


    descriptions_sheet = workbook.create_sheet("Descriptions")
    descriptions_sheet.append(["ASIN", "Content Type", "Content"])
    description_fields = (
        ("description", "Description"),
        ("product_description", "Product Description"),
        ("from_the_brand", "From the Brand"),
        ("customer_says", "Customer Says"),
        ("customers_say", "Customers Say"),
        ("top_review", "Top Review"),
        ("editorial_reviews", "Editorial Reviews"),
        ("about_the_author", "About the Author"),
        ("ingredients", "Ingredients"),
        ("safety_information", "Safety Information"),
        ("return_policy", "Return Policy"),
    )
    for record in records:
        asin = _record_asin(record)
        for field, label in description_fields:
            if record.get(field) not in (None, "", [], {}):
                descriptions_sheet.append([asin, label, _human_value(record.get(field))])
    _style_table(
        descriptions_sheet,
        widths={"ASIN": 16, "Content Type": 25, "Content": 120},
        wrap_headers={"Content"},
    )


    data_sheet = workbook.create_sheet("Amazon Data")
    fields = _ordered_fields(records)
    data_sheet.append(fields)
    for record in records:
        data_sheet.append([_excel_value(record.get(field)) for field in fields])
    _style_table(data_sheet, max_width=28)


    audit_sheet = workbook.create_sheet("Input Audit")
    audit_sheet.append(
        ["Source Row", "Original Value", "Normalized ASIN", "Status", "Message"]
    )
    for audit_row in audit_rows:
        values = asdict(audit_row)
        audit_sheet.append(
            [
                _excel_value(values["source_row"]), _excel_value(values["original_value"]),
                _excel_value(values["normalized_asin"]), _excel_value(values["status"]),
                _excel_value(values["message"]),
            ]
        )
    _style_table(
        audit_sheet,
        widths={
            "Source Row": 14, "Original Value": 45, "Normalized ASIN": 18,
            "Status": 16, "Message": 55,
        },
        wrap_headers={"Message"},
        number_formats={"Source Row": "0"},
    )


    error_sheet = workbook.create_sheet("Errors")
    error_fields = (
        "input_url", "input_asin", "asin", "error", "error_code", "warning",
        "warning_code",
    )
    error_sheet.append(
        [
            "Input URL", "Input ASIN", "Returned ASIN", "Error", "Error Code",
            "Warning", "Warning Code",
        ]
    )
    for record in records:
        if record.get("error") or record.get("warning"):
            error_sheet.append([_excel_value(record.get(field)) for field in error_fields])
    _style_table(
        error_sheet,
        widths={
            "Input URL": 65, "Input ASIN": 18, "Returned ASIN": 18,
            "Error": 65, "Error Code": 18, "Warning": 65, "Warning Code": 18,
        },
        wrap_headers={"Error", "Warning"},
        hyperlink_headers={"Input URL"},
    )


    run_sheet = workbook.create_sheet("Run Information")
    run_sheet.append(["Field", "Value"])
    for key, value in complete_metadata.items():
        run_sheet.append(
            [str(key).replace("_", " ").title(), _human_value(value)]
        )
    _style_table(
        run_sheet,
        widths={"Field": 32, "Value": 75},
        wrap_headers={"Value"},
    )


    workbook.active = 0
    workbook.save(workbook_path)
    workbook.close()
    return raw_json_path, workbook_path, metadata_path