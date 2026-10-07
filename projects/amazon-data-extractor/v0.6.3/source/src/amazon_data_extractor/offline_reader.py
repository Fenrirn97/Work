"""Read a previously generated Amazon extraction workbook for offline comparison."""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook

from .catalog_reader import clean_asin


@dataclass(frozen=True)
class OfflineExtraction:
    records: list[dict[str, Any]]
    requested_asins: list[str]
    data_sheet: str


class OfflineWorkbookError(ValueError):
    pass


def normalize_header(value: object) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(value or "").strip().casefold()).strip()


def _headers(worksheet) -> tuple[list[str], dict[str, int]]:
    values = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    headers = [str(value or "").strip() for value in values]
    indexes = {}
    for index, header in enumerate(headers):
        normalized = normalize_header(header)
        if normalized and normalized not in indexes:
            indexes[normalized] = index    return headers, indexes


def _requested_asins(workbook) -> list[str]:
    if "Input Audit" not in workbook.sheetnames:
        raise OfflineWorkbookError(
            "The extraction workbook has no 'Input Audit' sheet. "
            "Choose amazon_extraction.xlsx created by this program."
        )

    worksheet = workbook["Input Audit"]
    _header_values, columns = _headers(worksheet)
    asin_column = columns.get("normalized asin")
    status_column = columns.get("status")
    if asin_column is None or status_column is None:
        raise OfflineWorkbookError(
            "The 'Input Audit' sheet needs Normalized ASIN and Status columns."
        )

    requested = []
    seen = set()
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        status = str(row[status_column] or "").strip().casefold()
        asin = clean_asin(row[asin_column])
        if status == "accepted" and asin and asin not in seen:
            seen.add(asin)
            requested.append(asin)

    if not requested:
        raise OfflineWorkbookError(
            "The 'Input Audit' sheet contains no accepted ASINs."
        )
    return requested


def _records(workbook) -> tuple[list[dict[str, Any]], str]:
    if "Amazon Data" not in workbook.sheetnames:
        raise OfflineWorkbookError(
            "The extraction workbook has no 'Amazon Data' sheet. "
            "Choose amazon_extraction.xlsx created by this program."
        )

    worksheet = workbook["Amazon Data"]
    headers, _columns = _headers(worksheet)
    if not any(headers):
        raise OfflineWorkbookError("The 'Amazon Data' sheet has no headers.")

    records = []
    for row in worksheet.iter_rows(min_row=2, values_only=True):
        record = {
            header: row[index] if index < len(row) else None
            for index, header in enumerate(headers)
            if header
        }
        if any(value not in (None, "") for value in record.values()):
            records.append(record)

    return records, worksheet.title


def read_offline_extraction(path: str | Path) -> OfflineExtraction:
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise OfflineWorkbookError("The Amazon extraction must be an .xlsx workbook.")

    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except (OSError, ValueError) as exc:
        raise OfflineWorkbookError(f"Could not open the extraction workbook: {exc}") from exc

    try:
        requested_asins = _requested_asins(workbook)
        records, data_sheet = _records(workbook)
        return OfflineExtraction(records, requested_asins, data_sheet)
    finally:
        workbook.close()