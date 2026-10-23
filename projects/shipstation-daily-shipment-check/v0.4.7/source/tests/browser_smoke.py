"""Optional local UI test: python tests/browser_smoke.py (requires a Playwright browser)."""
import base64
import io
import os
import sys
import threading
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import make_server
from shipstation import parse_ship_date_range

calls=[]
def extractor(credentials,start,progress,ship_date_end=None):
    calls.append((start,ship_date_end))
    selected=parse_ship_date_range(start,ship_date_end)
    return dict(version='v1',shipDate=selected['iso'],shipDateStart=selected['iso'],shipDateEnd=selected['end'],
                displayDate=selected['display'],rangeDays=selected['days'],fileRange=selected['fileRange'],
                generatedAt='2026-10-23T14:00:00+00:00',rowCount=0,recordCount=0,missingTrackingCount=0,
                coverage='Synthetic UI test only',quantityNote='Test',rows=[],events=[],issues=[])
server=make_server(port=0,extractor=extractor)
thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
errors=[]
try:
    with sync_playwright() as playwright:
        browser=playwright.chromium.launch(headless=True)
        page=browser.new_page(viewport={'width':1440,'height':1100})
        page.on('pageerror',lambda err:errors.append(str(err)))
        page.goto(f'http://127.0.0.1:{server.server_port}')
        page.locator('#credentialsFile').set_input_files({'name':'test_credentials.txt','mimeType':'text/plain','buffer':b'API_KEY=test-key-1234\nAPI_SECRET=test-secret-1234'})
        page.locator('#shipDateStart').fill('2026-09-29')
        page.locator('#shipDateEnd').fill('2026-10-02')
        page.locator('#shipDateEnd').dispatch_event('change')
        assert '09/29/2026 through 10/02/2026' in page.locator('#rangeDescription').inner_text()
        with page.expect_download() as captured:
            page.locator('#reportButton').click()
        download=captured.value
        assert download.suggested_filename=='Shipped_2026-09-29_to_2026-10-02_Checked.xlsx'
        book=load_workbook(io.BytesIO(Path(download.path()).read_bytes()))
        assert book.sheetnames[0]=='2026-09-29 to 2026-10-02'
        assert calls==[('2026-09-29','2026-10-02')]
        page.locator('#shipDateStart').fill('2026-10-05')
        page.locator('#shipDateStart').dispatch_event('change')
        page.locator('#reportButton').click()
        assert not page.locator('#shipDateEnd').evaluate('(e)=>e.checkValidity()')
        assert len(calls)==1
        page.locator('#shipDateEnd').fill('2026-10-05')
        page.locator('#shipDateEnd').dispatch_event('change')
        with page.expect_download() as captured:
            page.locator('#reportButton').click()
        assert captured.value.suggested_filename=='Shipped_2026-10-05_Checked.xlsx'
        assert calls[-1]==('2026-10-05','2026-10-05')
        assert not errors, errors
        screenshot=os.environ.get('SSC_QA_SCREENSHOT')
        if screenshot:
            page.screenshot(path=screenshot,full_page=True)
        browser.close()
    print('PASS: date controls, inclusive range payload, range workbook download, invalid-range blocking, single-day download, no browser JS errors.')
finally:
    server.shutdown();server.server_close();thread.join()
