
"""Create the combined Excel workbook in memory."""

from __future__ import annotations

from datetime import date, datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from shipstation import ReportError

ITEM_COLUMNS = [
    ("Order Number", "orderNumber", 24),
    ("Item Name", "itemName", 42),
    ("SKU", "sku", 34),
    ("Quantity", "quantity", 10),
    ("Order Date", "orderDate", 14),
    ("Ship Date(s)", "shipDate", 25),
    ("Tracking Number(s)", "tracking", 28),
    ("Carrier (ShipStation)", "carrier", 20),
    ("Service", "service", 22),
    ("Record Source", "source", 30),
    ("Label Created", "labelCreated", 16),
    ("Label Printed (CSV)", "labelPrinted", 18),
    ("ShipStation Printed Status", "shipstationUiPrintStatus", 24),
    ("ShipStation Displayed Print Time", "shipstationDisplayedPrintTime", 30),
    ("ShipStation Label Printed Raw", "shipstationLabelPrintedRaw", 50),
    ("ShipStation UI Match", "shipstationUiPrintMatch", 28),
    ("ShipStation Capture Time", "shipstationUiCapturedAt", 28),
    ("Print Check", "printCheck", 32),
    ("Print Match", "printMatch", 24),
    ("Marketplace Notified", "marketplaceNotified", 20),
    ("Shipment Notification", "shipmentNotification", 20),
    ("Delivery Notification", "deliveryNotification", 20),
    ("Carrier Used", "normalizedCarrier", 16),
    ("Tracking Page", "trackingUrl", 48),
    ("Carrier Status", "carrierStatus", 24),
    ("Carrier Accepted?", "carrierAccepted", 18),
    ("Carrier Status Detail", "carrierStatusDescription", 48),
    ("First Carrier Scan", "firstCarrierScan", 24),
    ("Latest Carrier Scan", "latestCarrierScan", 24),
    ("Latest Scan Location", "latestCarrierLocation", 30),
    ("Estimated Delivery", "estimatedDelivery", 24),
    ("Actual Delivery", "actualDelivery", 24),
    ("Shipping Warning", "shippingWarning", 46),
    ("Carrier Check Error", "carrierCheckError", 42),
    ("Store / Source Code", "store", 24),
    ("Recipient", "recipient", 26),
    ("Unit Price", "unitPrice", 12),
    ("Line Total", "lineTotal", 13),
    ("Fulfillment Provider", "fulfillmentProvider", 22),
    ("Order ID", "orderId", 22),
    ("Item ID", "itemId", 22),
    ("Shipment IDs", "shipmentIds", 24),
    ("Record IDs", "recordIds", 24),
    ("Review Notes", "warnings", 54),
]

TRACKING_COLUMNS = [
    ("Order Number / Shipment Reference", "orderNumber", 30),
    ("Order ID / External Order ID", "orderId", 26),
    ("Ship Date", "shipDate", 14),
    ("Tracking Number(s)", "tracking", 28),
    ("Carrier", "carrier", 22),
    ("Service", "service", 22),
    ("Shipment ID", "shipmentId", 24),
    ("Record ID", "recordId", 24),
    ("Record Source", "source", 30),
    ("Label Created", "labelCreated", 16),
    ("Label Printed (CSV)", "labelPrinted", 18),
    ("ShipStation Printed Status", "shipstationUiPrintStatus", 24),
    ("ShipStation Displayed Print Time", "shipstationDisplayedPrintTime", 30),
    ("ShipStation UI Match", "shipstationUiPrintMatch", 28),
    ("Marketplace Notified", "marketplaceNotified", 20),
    ("Print Match", "printMatch", 24),
    ("Carrier Status", "carrierStatus", 24),
    ("Carrier Accepted?", "carrierAccepted", 18),
    ("Recipient", "recipient", 26),
]

PRINT_STATUS_COLUMNS = [
    ("Shipment #", "shipmentNumber", 22),
    ("Order #", "orderNumber", 26),
    ("Tracking # (raw CSV)", "trackingRaw", 28),
    ("Tracking # (safe text)", "tracking", 28),
    ("Service", "service", 22),
    ("Recipient", "recipient", 26),
    ("Ship Date", "shipDate", 16),
    ("Packing Slip Printed", "packingSlipPrinted", 20),
    ("Label Created", "labelCreated", 16),
    ("Label Printed (CSV)", "labelPrinted", 18),
    ("Marketplace Notified", "marketplaceNotified", 20),
    ("Shipment Notification", "shipmentNotification", 20),
    ("Delivery Notification", "deliveryNotification", 20),
]

UI_PRINT_CAPTURE_COLUMNS = [
    ("Shipment #", "shipmentNumber", 24),
    ("Order #", "orderNumber", 26),
    ("Tracking #", "tracking", 30),
    ("Printed Status", "printStatus", 24),
    ("Displayed Print Date / Time", "displayedPrintTime", 30),
    ("Raw Label Printed Cell", "rawLabelPrinted", 56),
    ("Captured At", "capturedAt", 28),
]

CARRIER_COLUMNS = [
    ("Tracking Number", "tracking", 30),
    ("Carrier", "carrier", 16),
    ("Tracking Page", "trackingUrl", 48),
    ("Status", "status", 24),
    ("Carrier Accepted?", "carrierAccepted", 18),
    ("Status Detail", "statusDescription", 52),
    ("First Carrier Scan", "firstCarrierScan", 24),
    ("Latest Scan", "latestScan", 24),
    ("Latest Location", "latestLocation", 30),
    ("Estimated Delivery", "estimatedDelivery", 24),
    ("Actual Delivery", "actualDelivery", 24),
    ("Checked At", "checkedAt", 24),
    ("Error", "error", 46),
]


def _excel_date(value):
    if not value:
        return None
    if isinstance(value, (date, datetime)):
        return value
    text = str(value).strip()
    try:
        if len(text) >= 10 and text[4] == "-":
            return datetime.fromisoformat(text.replace("Z", "+00:00")).replace(tzinfo=None)
        if len(text) == 10 and text[2] == "/":
            return datetime.strptime(text, "%m/%d/%Y")
    except ValueError:
        return value
    return value


def _cell_fill(key, value):
    text = str(value or "").upper()
    if key in ("shippingWarning", "carrierCheckError") and value:
        if "CRITICAL" in text or "ERROR" in text or "NO CARRIER ACCEPTANCE" in text:
            return PatternFill("solid", fgColor="F4CCCC")
        if "REVIEW" in text or "NOT CONFIGURED" in text or "UNSUPPORTED" in text:
            return PatternFill("solid", fgColor="FFF2CC")
    if key in ("printCheck", "labelPrinted", "shipstationUiPrintStatus"):
        if text in ("NOT PRINTED", "CREATED - NOT PRINTED") or text.startswith("REVIEW"):
            return PatternFill("solid", fgColor="F4CCCC" if "NOT PRINTED" in text else "FFF2CC")
        if text == "PRINTED":
            return PatternFill("solid", fgColor="D9EAD3")
        if text == "COULD NOT READ":
            return PatternFill("solid", fgColor="FFF2CC")
    if key == "carrierAccepted":
        if text == "NO":
            return PatternFill("solid", fgColor="F4CCCC")
        if text == "YES":
            return PatternFill("solid", fgColor="D9EAD3")
    if key == "carrierStatus":
        if "LABEL CREATED" in text:
            return PatternFill("solid", fgColor="FCE5CD")
        if text in ("DELIVERED", "IN TRANSIT", "CARRIER ACCEPTED", "OUT FOR DELIVERY"):
            return PatternFill("solid", fgColor="D9EAD3")
        if text in ("ERROR", "INVALID TRACKING", "EXCEPTION", "CANCELLED"):
            return PatternFill("solid", fgColor="F4CCCC")
        if text == "MANUAL CHECK REQUIRED":
            return PatternFill("solid", fgColor="FFF2CC")
    if key == "warnings" and value:
        return PatternFill("solid", fgColor="FFF2CC")
    return None


def table(sheet, columns, records):
    if len(records) > 1048575:
        raise ReportError("This report exceeds Excel's row limit.")
    sheet.append([header for header, _, _ in columns])
    for index, (_, _, width) in enumerate(columns, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    date_like_keys = {
        "orderDate",
        "shipDate",
        "firstCarrierScan",
        "latestCarrierScan",
        "estimatedDelivery",
        "actualDelivery",
        "checkedAt",
        "latestScan",
        "shipstationUiCapturedAt",
    }

    for row_index, record in enumerate(records, 2):
        for column_index, (_, key, _) in enumerate(columns, 1):
            value = record.get(key)
            converted = _excel_date(value) if key in date_like_keys else value
            if isinstance(converted, str) and len(converted) > 32767:
                converted = converted[:32764] + "..."
            cell = sheet.cell(row_index, column_index, converted)
            if isinstance(converted, str):
                cell.data_type = "s"
                cell.number_format = "@"
                if key == "trackingUrl" and converted.startswith("https://") and ";" not in converted:
                    cell.hyperlink = converted
                    cell.style = "Hyperlink"
            elif isinstance(converted, datetime):
                if key in ("orderDate", "shipDate"):
                    cell.number_format = "mm/dd/yyyy"
                else:
                    cell.number_format = "mm/dd/yyyy h:mm AM/PM"
            elif key in ("unitPrice", "lineTotal") and converted is not None:
                cell.number_format = "#,##0.00"
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            fill = _cell_fill(key, value)
            if fill:
                cell.fill = fill
        sheet.row_dimensions[row_index].height = 38

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 34
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="285F99")
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0


def build_excel(report):
    workbook = Workbook()
    workbook.properties.creator = "ShipStation Daily Shipment Check"
    workbook.properties.created = datetime.fromisoformat(report["generatedAt"]).replace(tzinfo=None)

    items = workbook.active
    start = report.get("shipDateStart", report["shipDate"])
    end = report.get("shipDateEnd", start)
    items.title = start if start == end else start + " to " + end
    table(items, ITEM_COLUMNS, report["rows"])
    table(workbook.create_sheet("Tracking Records"), TRACKING_COLUMNS, report["events"])
    table(workbook.create_sheet("Print Status CSV"), PRINT_STATUS_COLUMNS, report.get("printStatusRecords", []))
    table(workbook.create_sheet("ShipStation Print Capture"), UI_PRINT_CAPTURE_COLUMNS, report.get("shipstationPrintCapture", []))
    table(workbook.create_sheet("Carrier Tracking"), CARRIER_COLUMNS, report.get("carrierTracking", []))

    notes = [
        ("Program version", "0.4.7"),
        ("Ship Date range (inclusive)", report["displayDate"]),
        ("From Ship Date", start),
        ("Through Ship Date", end),
        ("Calendar days included", report.get("rangeDays", 1)),
        ("Range basis", "Ship Date recorded by ShipStation. This is not the order date, label creation date, or displayed print time. Both boundary dates are included."),
        ("Supporting inputs", "CSV and print-capture sheets preserve supplied inputs, which may include other dates. Matches are attached to extracted shipment records in the selected Ship Date range. Capture every ShipStation page in that same range for full print-time coverage."),
        ("Generated at (UTC)", report["generatedAt"]),
        ("ShipStation API mode", report["version"].upper()),
        ("Coverage", report["coverage"]),
        ("Order / shipment groups", report["recordCount"]),
        ("Item rows", report["rowCount"]),
        ("Groups missing tracking", report["missingTrackingCount"]),
        ("Print Status CSV selected", "Yes" if report.get("printCsvProvided") else "No"),
        ("Print Status CSV records", report.get("printCsvRecordCount", 0)),
        ("Rows matched to Print Status CSV", report.get("printCsvMatchedRows", 0)),
        ("Rows flagged NOT PRINTED", report.get("notPrintedRowCount", 0)),
        ("Unmatched CSV records", report.get("unmatchedPrintCsvCount", 0)),
        ("ShipStation UI print capture selected", "Yes" if report.get("uiPrintCaptureProvided") else "No"),
        ("ShipStation UI rows captured", report.get("uiPrintCaptureCount", 0)),
        ("Report rows matched to UI capture", report.get("uiPrintMatchedRows", 0)),
        ("Ambiguous UI print matches", report.get("uiPrintAmbiguousRows", 0)),
        ("Printed rows missing displayed print time", report.get("uiPrintedMissingTimeCount", 0)),
        ("Printed rows missing tracking", report.get("printedMissingTrackingCount", 0)),
        ("Carrier verification selected", "Yes" if report.get("carrierCheckEnabled") else "No"),
        ("Tracking numbers checked", report.get("carrierCheckCount", 0)),
        ("Carrier accepted", report.get("carrierAcceptedCount", 0)),
        ("No carrier acceptance", report.get("carrierNotAcceptedCount", 0)),
        ("Manual carrier checks", report.get("carrierManualCount", 0)),
        ("Carrier web check issues", report.get("carrierErrorCount", 0)),
        (
            "Meaning of shipped",
            "For this report, a package is considered actually shipped only after UPS, FedEx or USPS reports physical carrier acceptance. Label Created alone is NOT treated as shipped.",
        ),
        (
            "Print status",
            "The old v0.4.x CSV comparison is preserved. When the optional ShipStation Print Time Capture extension has sent rows from the Shipments page, the Shipments-grid Label Printed state becomes the primary print check for matched shipments.",
        ),
        (
            "Displayed print time",
            "ShipStation Displayed Print Time is copied from the Shipments -> Label Printed cell/tooltip captured in the user's normal logged-in browser. It is not inferred from label creation time, Ship Date, or Windows print history. The raw cell text is preserved in the workbook.",
        ),
        (
            "Print-time matching",
            "Shipments-grid observations match by exact tracking number. When API tracking is missing, only one exact shipment identifier can be used. An order number alone never transfers a print observation. Multiple-package item rows require review of Tracking Records.",
        ),
        (
            "Fulfillment records",
            "Fulfillment / marked shipped records are not automatically treated as missing print records because they may not represent a ShipStation-created label.",
        ),
        (
            "Matching order",
            "CSV observations match only by exact tracking number. Conflicting orders, missing tracking, and multiple-package item rows are not guessed. Tracking Records keeps each shipment separately.",
        ),
        (
            "Tracking numbers are text",
            "Tracking identifiers are always stored as text. Scientific-notation values from CSV are never used for a carrier website check because Excel may have destroyed their exact digits.",
        ),
        (
            "Supported carriers",
            "Automated carrier verification supports only UPS, FedEx and USPS. Unsupported carriers are not guessed.",
        ),
        (
            "Carrier website limitations",
            "If a carrier website blocks automated reading, asks for human verification, or changes its page, that tracking number is marked MANUAL CHECK REQUIRED. A clickable tracking-page link is kept for manual review.",
        ),
        ("Quantity meaning", report["quantityNote"]),
        (
            "Privacy",
            "ShipStation API credentials are used in memory for the current run and are not written into the workbook or application logs. Carrier checks use public tracking pages and require no carrier API credentials.",
        ),
    ] + [(f"Review {i}", issue) for i, issue in enumerate(report.get("issues", []), 1)]

    info = workbook.create_sheet("Read Me")
    table(
        info,
        [("Report detail", "detail", 34), ("Explanation", "value", 110)],
        [{"detail": name, "value": value} for name, value in notes],
    )
    for row in info.iter_rows(min_row=2):
        info.row_dimensions[row[0].row].height = 48

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
