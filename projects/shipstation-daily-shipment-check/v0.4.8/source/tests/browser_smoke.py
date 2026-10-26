"""Optional local UI test: python tests/browser_smoke.py (requires a Playwright browser)."""
import base64
import io
import json
import os
import sys
import threading
import tempfile
from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import sync_playwright, expect
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from app import make_server
from shipstation import parse_ship_date_range

calls=[]
credential_calls=[]
def json_location(path):
    return json.loads(path.read_text())['location']

def extractor(credentials,start,progress,ship_date_end=None):
    calls.append((start,ship_date_end))
    credential_calls.append(credentials)
    selected=parse_ship_date_range(start,ship_date_end)
    return dict(version='v1',shipDate=selected['iso'],shipDateStart=selected['iso'],shipDateEnd=selected['end'],
                displayDate=selected['display'],rangeDays=selected['days'],fileRange=selected['fileRange'],
                generatedAt='2026-10-23T14:00:00+00:00',rowCount=0,recordCount=0,missingTrackingCount=0,
                coverage='Synthetic UI test only',quantityNote='Test',rows=[],events=[],issues=[])
temporary=tempfile.TemporaryDirectory()
test_root=Path(temporary.name)
auto_key='API_KEY=synthetic-auto-1111\nAPI_SECRET=synthetic-secret-1111'
custom_key='API_KEY=synthetic-custom-2222\nAPI_SECRET=synthetic-secret-2222'
auto_file=test_root/'API key.txt'; auto_file.write_text(auto_key)
custom_file=test_root/'Other Folder'/'API key.txt'
custom_file.parent.mkdir(); custom_file.write_text(custom_key)
settings=test_root/'settings.json'
server=make_server(port=0,extractor=extractor,app_folder=test_root,settings_path=settings)
thread=threading.Thread(target=server.serve_forever,daemon=True);thread.start()
errors=[]
try:
    with sync_playwright() as playwright:
        browser=playwright.chromium.launch(headless=True)
        page=browser.new_page(viewport={'width':1440,'height':1100})
        page.on('pageerror',lambda err:errors.append(str(err)))
        report_requests=[]
        page.on('request',lambda r:report_requests.append(r.post_data_json) if r.url.endswith('/api/report') else None)
        page.goto(f'http://127.0.0.1:{server.server_port}')
        expect(page.locator('#credentialsStatus')).to_contain_text('Ready')
        assert 'synthetic' not in page.locator('body').inner_text()
        page.locator('#shipDateStart').fill('2026-09-29')
        page.locator('#shipDateEnd').fill('2026-10-02')
        page.locator('#shipDateEnd').dispatch_event('change')
        assert '09/29/2026 through 10/02/2026' in page.locator('#rangeDescription').inner_text()
        with page.expect_download() as captured:
            page.locator('#reportButton').click()
        download=captured.value
        assert download.suggested_filename=='Shipped_2026-09-29_to_2026-10-02_Checked.xlsx'
        book=load_workbook(io.BytesIO(Path(download.path()).read_bytes()))
        assert book.sheetnames[0]=='2026-09-29 to 2026-10-02'
        assert calls==[('2026-09-29','2026-10-02')]
        assert credential_calls==[auto_key]
        assert 'credentialsText' not in report_requests[-1]
        page.locator('#shipDateStart').fill('2026-10-05')
        page.locator('#shipDateStart').dispatch_event('change')
        page.locator('#reportButton').click()
        assert not page.locator('#shipDateEnd').evaluate('(e)=>e.checkValidity()')
        assert len(calls)==1
        page.locator('#shipDateEnd').fill('2026-10-05')
        page.locator('#shipDateEnd').dispatch_event('change')
        with page.expect_download() as captured:
            page.locator('#reportButton').click()
        assert captured.value.suggested_filename=='Shipped_2026-10-05_Checked.xlsx'
        assert calls[-1]==('2026-10-05','2026-10-05')
        page.locator('#credentialsOptions summary').click()
        page.locator('#credentialsPath').fill(str(custom_file.parent))
        # Submitting with a typed path also selects and remembers that location.
        with page.expect_download():
            page.locator('#reportButton').click()
        assert credential_calls[-1]==custom_key
        assert json_location(settings)==str(custom_file)
        page.reload()
        expect(page.locator('#credentialsLocation')).to_contain_text(str(custom_file))
        page.locator('#credentialsOptions summary').click()
        expect(page.locator('#credentialsPath')).to_have_value(str(custom_file))
        with page.expect_download():
            page.locator('#reportButton').click()
        assert credential_calls[-1]==custom_key
        # A selected upload overrides the saved location for this report only.
        manual=b'API_KEY=test-key-3333\nAPI_SECRET=test-secret-3333'
        page.locator('#credentialsFile').set_input_files({'name':'test_credentials.txt','mimeType':'text/plain','buffer':manual})
        with page.expect_download():
            page.locator('#reportButton').click()
        assert credential_calls[-1]==manual.decode()
        assert json_location(settings)==str(custom_file)
        page.locator('#resetCredentialsPath').click()
        expect(page.locator('#credentialsLocation')).to_contain_text(str(auto_file))
        assert page.locator('#credentialsFile').evaluate('(e)=>e.files.length')==0
        with page.expect_download():
            page.locator('#reportButton').click()
        assert credential_calls[-1]==auto_key
        assert 'credentialsText' not in report_requests[-1]
        # Bad supplied paths block extraction rather than using a different key.
        before=len(calls)
        page.locator('#credentialsPath').fill(str(test_root/'missing.txt'))
        page.locator('#reportButton').click()
        expect(page.locator('#statusMessage')).to_contain_text('not found')
        assert len(calls)==before
        page.locator('#resetCredentialsPath').click()
        expect(page.locator('#credentialsPath')).to_have_value('')
        assert 'synthetic-secret' not in page.locator('body').inner_text()
        assert not errors, errors
        screenshot=os.environ.get('SSC_QA_SCREENSHOT')
        if screenshot:
            page.screenshot(path=screenshot,full_page=True)
        browser.close()
    print('PASS: automatic key loading, custom path remembered, upload override, program-folder reset, missing-file blocking, date ranges, Excel downloads, no secrets in automatic report requests, no browser JS errors.')
finally:
    server.shutdown();server.server_close();thread.join()
    temporary.cleanup()
