import base64
import copy
import io
import json
import sys
import threading
import unittest
from pathlib import Path
from unittest.mock import patch
from urllib.request import Request, urlopen
from urllib.error import HTTPError
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app import make_server
from shipstation import (ReportError, parse_ship_date_range, ship_date_in_range,
                         get_v1_report, get_v2_report, extract_report)
from print_status import merge_print_status, parse_print_status_csv
from print_capture import normalize_capture_rows, merge_shipstation_print_capture, parse_label_print_cell


class FakeAPI:
    def __init__(self, data=None, details=None):
        self.data = data or {}; self.details = details or {}; self.calls=[]
        self.progress=lambda *_:None; self.interval=0
    def get_pages(self,path,key,**params):
        self.calls.append((path,params))
        yield copy.deepcopy(self.data.get(path,[]))
    def get(self,path,**params):
        self.calls.append((path,params))
        return copy.deepcopy(self.details.get(path,[] if path=='/stores' else {}))


def fixture_v1():
    def shipment(n,day,order_id=10):
        return dict(shipmentId=n,orderId=order_id,orderNumber='ORDER-A',shipDate=day,
                    trackingNumber='TRACK-'+str(n),carrierCode='ups',serviceCode='ups_ground',
                    shipTo={'name':'Test Recipient'},shipmentCost=1)
    rows=[shipment(1,'2026-09-29T00:00:00'), shipment(2,'2026-10-02T23:59:59.999'),
          shipment(3,'2026-09-28T23:59:59'),shipment(4,'2026-10-03T00:00:00')]
    rows.append(dict(rows[0])); rows.append(dict(rows[0],shipmentId=5,voided=True))
    order=dict(orderId=10,orderNumber='ORDER-A',orderStatus='shipped',shipDate='2026-10-02',
               orderDate='2026-08-01',items=[dict(name='Test item',sku='ITEM-A',quantity=2,unitPrice=10)],
               shipTo={'name':'Test Recipient'})
    return FakeAPI({'/shipments':rows,'/orders':[order]}, {'/stores':[]})


class DateRangeTests(unittest.TestCase):
    def test_inclusive_range_across_month_boundary(self):
        selected=parse_ship_date_range('09/29/2026','10/02/2026')
        self.assertEqual(selected['days'],4)
        self.assertEqual(selected['next'],'2026-10-03')
        for stamp in ['2026-09-29T00:00:00','2026-10-02T23:59:59.999']:
            self.assertTrue(ship_date_in_range(stamp,selected))
        for stamp in ['2026-09-28T23:59:59','2026-10-03T00:00:00','bad','']:
            self.assertFalse(ship_date_in_range(stamp,selected))
    def test_iso_picker_and_single_day_compatibility(self):
        self.assertEqual(parse_ship_date_range('2026-10-01')['display'],'10/01/2026')
        self.assertEqual(parse_ship_date_range('10/01/2026')['fileRange'],'2026-10-01')
    def test_leap_day_and_year_rollover(self):
        self.assertEqual(parse_ship_date_range('2024-02-28','2024-03-01')['days'],3)
        self.assertEqual(parse_ship_date_range('2026-12-31','2027-01-01')['next'],'2027-01-02')
    def test_reversed_invalid_and_out_of_bounds_dates(self):
        for start,end in [('2026-10-04','2026-10-01'),('2026-02-29','2026-03-01'),('1999-01-01',None)]:
            with self.assertRaises(ReportError): parse_ship_date_range(start,end)
    def test_v1_filters_boundaries_and_does_not_repeat_order_quantity(self):
        api=fixture_v1(); report=get_v1_report(api,parse_ship_date_range('2026-09-29','2026-10-02'))
        self.assertEqual(len(report['events']),2)
        self.assertEqual(len(report['rows']),1)
        self.assertEqual(report['rows'][0]['quantity'],2)
        self.assertEqual(report['rows'][0]['shipDate'],'2026-09-29; 2026-10-02')
        self.assertEqual({e['shipDate'] for e in report['events']},{'2026-09-29','2026-10-02'})
        query=next(p for path,p in api.calls if path=='/shipments')
        self.assertEqual(query,dict(shipDateStart='2026-09-29 00:00:00',shipDateEnd='2026-10-03 00:00:00',includeShipmentItems='true'))
        self.assertEqual(sum(path=='/orders' for path,_ in api.calls),1)
    def test_v2_scans_labels_once_and_filters_ship_date_not_create_date(self):
        labels=[dict(label_id='L1',shipment_id='se-1',status='completed',ship_date='2026-09-29T00:00:00Z',created_at='2026-08-01',tracking_number='T1'),
                dict(label_id='L2',shipment_id='se-1',status='completed',ship_date='2026-10-02T23:59:59Z',tracking_number='T2'),
                dict(label_id='L3',shipment_id='se-2',status='completed',ship_date='2026-10-03T00:00:00Z',tracking_number='T3')]
        api=FakeAPI({'/labels':labels},{'/shipments/se-1':dict(items=[dict(name='Item',sku='A',quantity=2)])})
        report=get_v2_report(api,parse_ship_date_range('2026-09-29','2026-10-02'))
        self.assertEqual(len(report['events']),2);self.assertEqual(len(report['rows']),1)
        self.assertEqual(sum(path=='/labels' for path,_ in api.calls),1)
        query=next(p for path,p in api.calls if path=='/fulfillments')
        self.assertEqual(query['ship_date_end'],'2026-10-03T00:00:00Z')
    def test_extract_report_metadata(self):
        api=fixture_v1(); api.version='v1'; api.headers={}; api.session=type('Session',(),{'close':lambda self:None})()
        with patch('shipstation.ShipStationAPI',return_value=api):
            report=extract_report('API_KEY=test-key-1234\nAPI_SECRET=test-secret-1234','2026-09-29',ship_date_end='2026-10-02')
        self.assertEqual(report['fileRange'],'2026-09-29_to_2026-10-02')
        self.assertEqual(report['rangeDays'],4)
        self.assertEqual(report['recordCount'],1)


class MatchRegressionTests(unittest.TestCase):
    def report(self,tracking):
        return {'rows':[dict(orderNumber='ORDER-A',orderId='10',shipmentIds='1',tracking=tracking,source='ShipStation label',warnings='',labelPrinted='')],'events':[]}
    def test_csv_different_package_same_order_is_not_printed_match(self):
        report=self.report('TRACK-B')
        data='Order #,Tracking #,Label Created,Label Printed\nORDER-A,TRACK-A,Created,Printed\n'
        merge_print_status(report,parse_print_status_csv(data))
        self.assertNotEqual(report['rows'][0]['printCheck'],'PRINTED')
    def test_capture_different_package_same_order_is_not_match(self):
        report=self.report('TRACK-B')
        captured=normalize_capture_rows([dict(shipmentNumber='1',orderNumber='ORDER-A',tracking='TRACK-A',labelPrintedRaw='Printed 09/29/2026 10:30')])
        merge_shipstation_print_capture(report,captured)
        self.assertEqual(report['uiPrintMatchedRows'],0)
    def test_multiple_package_group_does_not_inherit_one_time(self):
        report=self.report('TRACK-A; TRACK-B')
        captured=normalize_capture_rows([dict(orderNumber='ORDER-A',tracking='TRACK-A',labelPrintedRaw='Printed 09/29/2026 10:30')])
        merge_shipstation_print_capture(report,captured)
        self.assertEqual(report['uiPrintMatchedRows'],0)
    def test_compact_display_timestamp(self):
        for raw in ['Printed09/17/2026•11:36','Printed 09/17/2026·11:36']:
            parsed=parse_label_print_cell(raw)
            self.assertEqual(parsed['printStatus'],'Printed')
            self.assertEqual(parsed['displayedPrintTime'],'09/17/2026 11:36')


class HTTPRangeTests(unittest.TestCase):
    def setUp(self):
        self.received=[]
        def extractor(credentials,start,progress,ship_date_end=None):
            self.received.append((start,ship_date_end))
            selected=parse_ship_date_range(start,ship_date_end)
            return dict(version='v1',shipDate=selected['iso'],shipDateStart=selected['iso'],shipDateEnd=selected['end'],
                        displayDate=selected['display'],rangeDays=selected['days'],fileRange=selected['fileRange'],
                        generatedAt='2026-10-23T14:00:00+00:00',rowCount=0,recordCount=0,missingTrackingCount=0,
                        coverage='Local test fixture',quantityNote='Test',rows=[],events=[],issues=[])
        self.server=make_server(port=0,extractor=extractor)
        self.thread=threading.Thread(target=self.server.serve_forever,daemon=True);self.thread.start()
        self.origin=f'http://127.0.0.1:{self.server.server_port}'
    def tearDown(self):
        self.server.shutdown();self.server.server_close();self.thread.join()
    def request_report(self,**values):
        payload=dict(credentialsText='API_KEY=test-key-1234\nAPI_SECRET=test-secret-1234',checkCarriers=False,**values)
        request=Request(self.origin+'/api/report',data=json.dumps(payload).encode(),headers={'Content-Type':'application/json','Origin':self.origin})
        return urlopen(request).read().decode()
    def test_range_reaches_extractor_and_excel(self):
        messages=[json.loads(line) for line in self.request_report(shipDateStart='2026-09-29',shipDateEnd='2026-10-02').splitlines()]
        result=next(m['report'] for m in messages if m['type']=='complete')
        self.assertEqual(self.received,[('2026-09-29','2026-10-02')])
        self.assertEqual(result['fileRange'],'2026-09-29_to_2026-10-02')
        book=load_workbook(io.BytesIO(base64.b64decode(result['fileBase64'])))
        self.assertEqual(book.sheetnames[0],'2026-09-29 to 2026-10-02')
        notes=dict(book['Read Me'].iter_rows(min_row=2,values_only=True))
        self.assertEqual(notes['Through Ship Date'],'2026-10-02')
    def test_reversed_range_rejected_before_api(self):
        with self.assertRaises(HTTPError) as caught:
            self.request_report(shipDateStart='2026-10-03',shipDateEnd='2026-10-01')
        self.assertEqual(caught.exception.code,400)
        self.assertEqual(self.received,[])
    def test_legacy_single_day_request_still_works(self):
        text=self.request_report(shipDate='09/29/2026')
        self.assertIn('"type": "complete"',text)
        self.assertEqual(self.received,[('09/29/2026',None)])


if __name__=='__main__': unittest.main()
