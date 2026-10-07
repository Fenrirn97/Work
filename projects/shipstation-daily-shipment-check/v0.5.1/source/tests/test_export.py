
import base64
import copy
from datetime import date, datetime
from http.client import HTTPConnection
from io import BytesIO
import json
from pathlib import Path
import sys
import threading
import unittest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook
from requests.structures import CaseInsensitiveDict
from app import make_server
from excel_export import build_excel
from shipstation import (
    ShipStationAPI,
    ReportError,
    extract_report,
    parse_credentials,
    parse_ship_date,
)

V1 = "API_KEY=test_key_123456\nAPI_SECRET=test_secret_123456"
V2 = "V2_API_KEY=test_key_123456"
DAY = "08/27/2026"
FAST = {"interval": 0, "sleep": lambda seconds: None}


class Response:
    def __init__(self, data, status=200, headers=None):
        self.data, self.status_code = data, status
        self.headers = CaseInsensitiveDict(headers or {})

    def json(self):
        if isinstance(self.data, Exception):
            raise self.data
        return copy.deepcopy(self.data)


def envelope(key, rows, page=1, pages=1, total=None):
    return {key: rows, "page": page, "pages": pages, "total": len(rows) if total is None else total}


class Session:
    def __init__(self, responder):
        self.responder, self.calls, self.closed = responder, [], False

    def get(self, url, **kwargs):
        self.calls.append((url, copy.deepcopy({k: v for k, v in kwargs.items() if k != "auth"})))
        return self.responder(url, kwargs)

    def close(self):
        self.closed = True


def v1_session(missing=False, multiple=False, extra=False, ship_day="2026-08-27"):
    orders = [
        {
            "orderId": i + 1,
            "orderNumber": "00000000000000000" + str(i + 1),
            "orderStatus": "shipped",
            "orderDate": "2026-08-18T11:00:00" if i == 6 else "2026-08-26T11:00:00",
            "shipDate": ship_day + "T00:00:00",
            "shipTo": {"name": "Test Recipient"},
            "advancedOptions": {"storeId": 42},
            "items": [
                {
                    "orderItemId": 100 + i,
                    "name": "Test item",
                    "sku": "SAME-WATCH-SKU" if i >= 3 else "SKU-" + str(i),
                    "quantity": 1,
                    "unitPrice": 12.5,
                }
            ],
        }
        for i in range(7)
    ]
    events = [
        {
            "orderId": i + 1,
            "orderNumber": order["orderNumber"],
            "shipDate": ship_day + "T00:00:00",
            "trackingNumber": "" if missing and i == 6 else "000000000000000000000" + str(i),
            "carrierCode": "ups" if i < 2 else "amazon",
            **(
                {"shipmentId": "label-" + str(i)}
                if i < 2
                else {"fulfillmentId": "external-" + str(i)}
            ),
        }
        for i, order in enumerate(orders)
    ]
    if multiple:
        events.append(
            {**events[0], "shipmentId": "extra-package", "trackingNumber": "SECOND-PACKAGE"}
        )
    if extra:
        orders.append({**orders[0], "orderId": 8, "orderNumber": "missing-event-order"})

    def respond(url, kwargs):
        path = url.removeprefix("https://ssapi.shipstation.com")
        if path.startswith("/orders/"):
            order_id = int(path.rsplit("/", 1)[-1])
            order = next((order for order in orders if order["orderId"] == order_id), None)
            return Response(order or {}, 200 if order else 404)
        key = url.rsplit("/", 1)[-1]
        if key == "fulfillments":
            return Response(envelope(key, [e for e in events if "fulfillmentId" in e]))
        if key == "shipments":
            return Response(envelope(key, [e for e in events if "shipmentId" in e]))
        if key == "stores":
            return Response([{"storeId": 42, "storeName": "Test Store"}])
        if key == "orders":
            page = kwargs["params"]["page"]
            return Response(
                envelope(key, orders[:3] if page == 1 else orders[3:], page, 2, len(orders))
            )
        raise AssertionError("Unexpected fixture path")

    return Session(respond)


def v2_session(missing_detail=False):
    def respond(url, kwargs):
        path = url.removeprefix("https://api.shipstation.com/v2")
        if path == "/fulfillments":
            return Response(
                envelope(
                    "fulfillments",
                    [
                        {
                            "fulfillment_id": "f1",
                            "shipment_id": "s1",
                            "shipment_number": "EXTERNAL-ORDER",
                            "ship_date": "2026-08-27T00:00:00Z",
                            "tracking_number": "AMAZON-TRACKING",
                            "fulfillment_carrier_friendly_name": "Amazon Logistics",
                        }
                    ],
                )
            )
        if path == "/labels":
            page = kwargs["params"]["page"]
            records = (
                [
                    {
                        "label_id": "l1",
                        "shipment_id": "s2",
                        "ship_date": "2026-08-27T00:00:00Z",
                        "created_at": "2026-08-25T11:00:00Z",
                        "tracking_number": "FIRST",
                        "status": "completed",
                    }
                ]
                if page == 1
                else [
                    {
                        "label_id": "l2",
                        "shipment_id": "s2",
                        "ship_date": "2026-08-27T00:00:00Z",
                        "tracking_number": "SECOND",
                        "status": "completed",
                    },
                    {"label_id": "outside", "ship_date": "2026-08-26T00:00:00Z"},
                ]
            )
            return Response(envelope("labels", records, page, 2, 3))
        if path.startswith("/shipments/"):
            if missing_detail and path.endswith("/s1"):
                return Response({}, 404)
            return Response(
                {
                    "shipment_number": "ORDER-" + path.split("/")[-1],
                    "created_at": "2026-08-20T00:00:00Z",
                    "items": [
                        {"sku": "TEST-SKU", "name": "Test item", "quantity": 2, "unit_price": 15}
                    ],
                }
            )
        raise AssertionError("Unexpected fixture path")

    return Session(respond)


class ExportTests(unittest.TestCase):
    def test_credentials_and_dates(self):
        self.assertEqual(parse_credentials(V1)["version"], "v1")
        self.assertEqual(parse_credentials("\ufeff# key\n" + V2)["version"], "v2")
        self.assertEqual(parse_credentials("Abcd1234567890==")["key"], "Abcd1234567890==")
        self.assertEqual(
            parse_credentials("API Key: abc1234567890\nAPI Secret: abc4567890123")["version"], "v1"
        )
        self.assertEqual(parse_ship_date("02/29/2024")["next"], "2024-03-01")
        self.assertEqual(parse_ship_date("12/31/2026")["next"], "2027-01-01")
        for value in (
            "",
            "YOUR_API_KEY_HERE",
            "V1_API_KEY=test_123456",
            V1 + "\n" + V2,
            "API_SECRET=test_123456",
        ):
            with self.subTest(value=value), self.assertRaises(ReportError):
                parse_credentials(value)
        for value in ("8/27/2026", "02/29/2026", "2026-08-27", "13/01/2026"):
            with self.subTest(value=value), self.assertRaises(ReportError):
                parse_ship_date(value)

    def test_seven_orders_include_five_external_fulfillments(self):
        session = v1_session()
        progress = []
        report = extract_report(V1, DAY, session=session, progress=progress.append, **FAST)
        self.assertEqual(
            (report["recordCount"], report["rowCount"], len(report["events"])), (7, 7, 7)
        )
        self.assertEqual(sum("Fulfillment" in r["source"] for r in report["rows"]), 5)
        self.assertEqual(sum(r["sku"] == "SAME-WATCH-SKU" for r in report["rows"]), 4)
        self.assertEqual(report["rows"][6]["orderDate"], "2026-08-18")
        self.assertEqual(report["rows"][0]["store"], "Test Store")
        self.assertTrue(session.closed)
        page_updates = [update for update in progress if isinstance(update, dict)]
        self.assertTrue(
            any(
                update["section"] == "Orders"
                and update["currentPage"] == 1
                and update["totalPages"] == 2
                for update in page_updates
            )
        )
        self.assertTrue(all(update["estimatedSeconds"] >= 0 for update in page_updates))
        for url, kwargs in session.calls:
            self.assertTrue(url.startswith("https://ssapi.shipstation.com/"))
            self.assertFalse(kwargs["allow_redirects"])
            self.assertEqual(kwargs["timeout"], (10, 45))
            if url.endswith("/orders"):
                self.assertEqual(kwargs["params"]["orderStatus"], "shipped")
                self.assertNotIn("orderDateStart", kwargs["params"])
                self.assertNotIn("createDateStart", kwargs["params"])

    def test_today_uses_only_linked_order_details(self):
        today = date.today()
        session = v1_session(ship_day=today.isoformat())
        progress = []
        report = extract_report(
            V1, today.strftime("%m/%d/%Y"), session=session, progress=progress.append, **FAST
        )
        self.assertEqual((report["recordCount"], report["rowCount"]), (7, 7))
        self.assertEqual(sum(url.endswith("/orders") for url, _ in session.calls), 0)
        self.assertEqual(sum("/orders/" in url for url, _ in session.calls), 7)
        self.assertIn("Today's linked orders", report["coverage"])
        self.assertTrue(
            any(
                isinstance(update, dict) and update["section"] == "Today's order details"
                for update in progress
            )
        )

    def test_missing_tracking_and_events_are_not_dropped(self):
        report = extract_report(V1, DAY, session=v1_session(missing=True, extra=True), **FAST)
        self.assertEqual(report["recordCount"], 8)
        self.assertEqual(report["missingTrackingCount"], 2)
        self.assertIn("No matching daily", report["rows"][-1]["warnings"])

    def test_multiple_packages_do_not_duplicate_quantity(self):
        report = extract_report(V1, DAY, session=v1_session(multiple=True), **FAST)
        self.assertEqual(report["rowCount"], 7)
        self.assertEqual(len(report["events"]), 8)
        self.assertEqual(report["rows"][0]["quantity"], 1)
        self.assertIn("SECOND-PACKAGE", report["rows"][0]["tracking"])

    def test_windows_print_job_is_matched_without_printing(self):
        def fake_print_history(selected_date):
            self.assertEqual(selected_date, "2026-08-27")
            return {
                "available": True,
                "note": "Synthetic Windows print history used by this test.",
                "jobs": [
                    {
                        "recordId": 30701,
                        "printedAt": "2026-08-27T10:30:00-04:00",
                        "documentName": "Walmart_Label_000000000000000001.pdf",
                        "userName": "Test User",
                        "computerName": "PACKING-PC",
                        "printerName": "Zebra Label Printer",
                        "pages": 1,
                    },
                    {
                        "recordId": 30702,
                        "printedAt": "2026-08-27T11:00:00-04:00",
                        "documentName": "Unrelated_Document.pdf",
                        "userName": "Test User",
                        "computerName": "PACKING-PC",
                        "printerName": "Office Printer",
                        "pages": 2,
                    },
                ],
            }

        report = extract_report(
            V1,
            DAY,
            session=v1_session(),
            print_history_reader=fake_print_history,
            **FAST,
        )
        printed_row = next(row for row in report["rows"] if row["orderNumber"].endswith("1"))

        self.assertEqual(printed_row["printStatus"], "Printed - Windows record found")
        self.assertEqual(printed_row["printerName"], "Zebra Label Printer")
        self.assertEqual(printed_row["printMatch"], "Order number")
        self.assertEqual(report["printJobCount"], 1)
        self.assertEqual(report["printJobs"][0]["matchedOrderNumbers"], "000000000000000001")
        self.assertNotIn("Unrelated_Document.pdf", str(report["printJobs"]))

    def test_void_return_and_other_day_exclusion(self):
        original = v1_session()

        def respond(url, kwargs):
            response = original.responder(url, kwargs)
            if url.endswith("/shipments"):
                first = response.data["shipments"][0]
                response.data["shipments"].extend(
                    [
                        {**first, "shipmentId": "void", "voided": True},
                        {**first, "shipmentId": "return", "isReturnLabel": True},
                        {**first, "shipmentId": "tomorrow", "shipDate": "2026-08-28T00:00:00"},
                    ]
                )
            return response

        report = extract_report(V1, DAY, session=Session(respond), **FAST)
        self.assertEqual(len(report["events"]), 7)

    def test_v2_older_labels_external_tracking_and_missing_order_date(self):
        session = v2_session()
        report = extract_report(V2, DAY, session=session, **FAST)
        self.assertEqual((report["rowCount"], len(report["events"])), (2, 3))
        self.assertEqual(report["rows"][0]["tracking"], "AMAZON-TRACKING")
        self.assertEqual(report["rows"][0]["carrier"], "Amazon Logistics")
        self.assertEqual(report["rows"][1]["tracking"], "FIRST; SECOND")
        self.assertEqual(report["rows"][1]["quantity"], 2)
        self.assertEqual(report["rows"][0]["orderDate"], "")
        self.assertIn("Order date unavailable", report["rows"][0]["warnings"])
        for url, kwargs in session.calls:
            if url.endswith("/labels"):
                self.assertNotIn("created_at_start", kwargs["params"])
        self.assertEqual(sum(url.endswith("/shipments/s2") for url, _ in session.calls), 1)

    def test_v2_missing_detail_preserves_tracking(self):
        report = extract_report(V2, DAY, session=v2_session(missing_detail=True), **FAST)
        row = next(r for r in report["rows"] if r["tracking"] == "AMAZON-TRACKING")
        self.assertIsNone(row["quantity"])
        self.assertIn("Shipment detail not found", row["warnings"])
        self.assertIn("SKU unavailable", row["warnings"])

    def test_auth_errors_redact_body_and_do_not_retry(self):
        session = Session(lambda url, kwargs: Response("PRIVATE CREDENTIAL TEXT", 401))
        with self.assertRaises(ReportError) as raised:
            extract_report(V1, DAY, session=session, **FAST)
        self.assertEqual(raised.exception.code, "AUTH")
        self.assertNotIn("PRIVATE", str(raised.exception))
        self.assertEqual(len(session.calls), 1)

    def test_retry_and_partial_failure(self):
        count = 0

        def respond(url, kwargs):
            nonlocal count
            count += 1
            return Response({}, 429, {"retry-after": "0"}) if count == 1 else Response({"ok": True})

        self.assertTrue(
            ShipStationAPI(parse_credentials(V2), session=Session(respond), **FAST).get("/labels")[
                "ok"
            ]
        )
        with self.assertRaises(ReportError) as raised:
            extract_report(V2, DAY, session=Session(lambda url, kwargs: Response({}, 503)), **FAST)
        self.assertEqual(raised.exception.code, "RETRY_LIMIT")

    def test_pagination_failures(self):
        cases = [
            ({"labels": [{"label_id": "same"}], "pages": 2}, "PAGINATION"),
            ({"labels": [], "total": 3, "pages": 2}, "PAGINATION"),
            ({"unexpected": []}, "BAD_RESPONSE"),
        ]
        for data, code in cases:
            client = ShipStationAPI(
                parse_credentials(V2), session=Session(lambda url, kwargs: Response(data)), **FAST
            )
            with self.subTest(code=code), self.assertRaises(ReportError) as raised:
                list(client.get_pages("/labels", "labels"))
            self.assertEqual(raised.exception.code, code)

    def test_excel_round_trip_strings_dates_prices_and_missing_fields(self):
        report = extract_report(V1, DAY, session=v1_session(), **FAST)
        report["rows"][0]["sku"] = '=HYPERLINK("https://example.invalid","test")'
        book = load_workbook(BytesIO(build_excel(report)))
        self.assertEqual(
            book.sheetnames, ["2026-08-27", "Tracking Records", "Print History", "Read Me"]
        )
        sheet = book.worksheets[0]
        self.assertEqual(sheet.max_row, 8)
        self.assertEqual(sheet["A2"].value, "000000000000000001")
        self.assertEqual(sheet["G2"].value, "0000000000000000000000")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertEqual(sheet["E8"].value.date().isoformat(), "2026-08-18")
        self.assertEqual(sheet["F2"].number_format, "mm/dd/yyyy")
        self.assertEqual(sheet["J2"].value, 12.5)
        self.assertEqual(sheet.freeze_panes, "A2")
        self.assertNotIn("test_secret", str(list(book["Read Me"].values)))
        v2 = extract_report(V2, DAY, session=v2_session(), **FAST)
        self.assertIsNone(load_workbook(BytesIO(build_excel(v2))).worksheets[0]["E2"].value)

    def test_empty_export_is_valid_excel(self):
        session = Session(lambda url, kwargs: Response(envelope(url.rsplit("/", 1)[-1], [])))
        report = extract_report(V2, DAY, session=session, **FAST)
        self.assertEqual(report["rowCount"], 0)
        self.assertEqual(load_workbook(BytesIO(build_excel(report))).worksheets[0].max_row, 1)

    def test_local_server_returns_python_generated_excel(self):
        public = ROOT / "public"
        if not public.exists():
            public = ROOT.parent / "standalone" / "public"
        calls = []

        def extractor(key, day, progress):
            calls.append(day)
            return extract_report(key, day, progress=progress, session=v1_session(), **FAST)

        server = make_server(0, root=public, extractor=extractor)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        origin = f"http://127.0.0.1:{server.server_port}"

        def request(method, path, data=None, headers=None):
            connection = HTTPConnection("127.0.0.1", server.server_port, timeout=5)
            connection.request(
                method,
                path,
                body=json.dumps(data) if data is not None else None,
                headers=headers or {},
            )
            response = connection.getresponse()
            result = response.status, response.read(), dict(response.getheaders())
            connection.close()
            return result

        try:
            status, page, _ = request("GET", "/")
            self.assertEqual(status, 200)
            self.assertIn(b"ShipStation Excel Export", page)
            self.assertEqual(request("GET", "/app.py")[0], 404)
            self.assertEqual(request("GET", "/", headers={"Host": "evil.invalid"})[0], 403)
            headers = {"Origin": origin, "Content-Type": "application/json"}
            self.assertEqual(
                request("POST", "/api/report", {}, {**headers, "Origin": "https://evil.invalid"})[
                    0
                ],
                403,
            )
            self.assertEqual(
                request("POST", "/api/report", {"credentialsText": "a" * 25000}, headers)[0], 400
            )
            self.assertEqual(
                request(
                    "POST",
                    "/api/report",
                    {"credentialsText": V1, "shipDate": "02/30/2026"},
                    headers,
                )[0],
                400,
            )
            self.assertEqual(calls, [])
            status, body, response_headers = request(
                "POST", "/api/report", {"credentialsText": V1, "shipDate": DAY}, headers
            )
            self.assertEqual(status, 200)
            self.assertEqual(response_headers["Cache-Control"], "no-store")
            events = [json.loads(line) for line in body.splitlines()]
            self.assertEqual(events[-1]["type"], "complete")
            self.assertNotIn(b"test_secret", body)
            book = load_workbook(BytesIO(base64.b64decode(events[-1]["report"]["fileBase64"])))
            self.assertEqual(book.worksheets[0].max_row, 8)
            self.assertEqual(calls, [DAY])
            status, body, _ = request("POST", "/api/exit", {}, headers)
            self.assertEqual(status, 200)
            self.assertTrue(json.loads(body)["stopping"])
            thread.join(2)
            self.assertFalse(thread.is_alive())
        finally:
            server.shutdown()
            server.server_close()
            thread.join(2)


if __name__ == "__main__":
    unittest.main()
